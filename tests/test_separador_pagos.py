"""Datos sintéticos; ningún libro real se utiliza como fixture."""
from datetime import date
from copy import copy
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

import separador_pagos as p

try:
    from streamlit.testing.v1 import AppTest
except ImportError:
    AppTest = None


HEADERS = ['SERIAL', 'CONCATENAR', 'FECHA DE PAGO', 'ESTATUS', 'OBSERVACION',
           'FECHA DE ARCHIVO', 'CANAL', 'CONCATENAR.1', 'AFILIADO', 'TERMINAL',
           'FECHA', 'VENDEDOR', 'EQUIPO', 'SERIAL.1', 'BANCO', 'NUMERO DE CUENTA',
           'GARANTIA DE POS', 'TOTAL VENTA ESTIMADO $$', 'ESTATUS CXC', 'FECHA DE PAGO',
           p.JORNADA, *p.PARES['BANCOS'], *p.PARES['REGIONALES'], *p.PARES['AGENTES AUTORIZADOS'],
           'MONTO TOTAL A PAGAR $', 'CON TX', 'MONTO TX AGOSTO',
           'REGISTRO DE OPERADORES', 'REGISTROS DE OPERADORES ACCESS', '__ROW_ID']


def fila(**cambios):
    r = {'CANAL': 'REGION ORIENTE', 'EQUIPO': 'CASTLE', 'BANCO': 'BANCO FICTICIO',
         'VENDEDOR': 'Persona Ficticia', p.JORNADA: 'OFICINA FICTICIA', 'ESTATUS': 'PAGADO',
         'OBSERVACION': 'HISTORICA', 'FECHA DE PAGO': date(2026, 1, 5), 'FECHA': date(2026, 1, 2),
         'VENDEDOR AGENTE AUTORIZADO': 'REGION ORIENTE', 'MONTO COMISION AGENTE AUTORIZADO $': 20,
         'MONTO TOTAL A PAGAR $': 38.4, 'CON TX': 'CON_TX', 'MONTO TX AGOSTO': 1501,
         'SERIAL': '123456789012', 'SERIAL.1': '123456789012', 'AFILIADO': '00001234',
         'REGISTRO DE OPERADORES': 123456, '__ROW_ID': 999}
    r.update(cambios)
    return r


def libro(filas, headers=None, hoja='VENTAS'):
    headers = headers or HEADERS
    w = Workbook()
    s = w.active
    s.title = hoja
    w.create_sheet('NO EXPORTAR')
    s.append(headers)
    for c in s[1]:
        c.font = Font(name='Calibri', size=9, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='123456')
        c.alignment = Alignment(wrap_text=True)
        c.border = Border(bottom=Side(style='thin', color='112233'))
    s.row_dimensions[1].height = 38
    for n, r in enumerate(filas, 2):
        s.append([r.get(h) for h in headers])
        s.row_dimensions[n].height = 25 + n
        for c in s[n]:
            c.font = Font(name='Calibri', size=9)
            c.fill = PatternFill('solid', fgColor='E0E0E0')
            c.alignment = Alignment(horizontal='center')
            c.border = Border(bottom=Side(style='thin', color='112233'))
            if isinstance(c.value, date):
                c.number_format = 'dd/mm/yyyy'
            elif isinstance(c.value, (int, float)):
                c.number_format = '#,##0.00'
    for c in s[1]:
        s.column_dimensions[c.column_letter].width = 20 + c.column / 10
    s.auto_filter.ref = s.dimensions
    b = BytesIO()
    w.save(b)
    w.close()
    return b.getvalue()


def salidas(datos):
    r = p.separar_pagos(datos, date(2026, 9, 16))
    with ZipFile(BytesIO(r.contenido_zip)) as z:
        libros = {n: load_workbook(BytesIO(z.read(n)), data_only=False) for n in z.namelist()}
    return r, libros


def salida(libros, nombre):
    return libros[f'Pago de comisiones {nombre} 18-09-2026.xlsx']['VENTAS']


def valores(s):
    headers = [c.value for c in s[1]]
    return [dict(zip(headers, row)) for row in s.iter_rows(min_row=2, max_row=int(s.auto_filter.ref.split(':')[1].lstrip('ABCDEFGHIJKLMNOPQRSTUVWXYZ')), values_only=True)]


class SeparadorPagosTests(unittest.TestCase):
    @unittest.skipIf(AppTest is None, 'Requiere Streamlit instalado para probar la interfaz.')
    def test_streamlit_navegacion_generacion_y_cambio_de_archivo(self):
        at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'), default_timeout=30).run()
        self.assertFalse(at.exception)
        self.assertEqual([x.proto.label for x in at.get('page_link')], ['Procesamiento de Comisiones', 'Separador de Pagos'])
        with patch('streamlit.file_uploader', return_value=BytesIO(libro([fila()]))):
            at.switch_page('pages/2_Separador_de_Pagos.py').run()
            at.button[0].click().run()
            self.assertFalse(at.exception)
            self.assertEqual(len(at.dataframe), 3)
            self.assertEqual(at.metric[0].value, '1')
            self.assertEqual(len(at.get('download_button')), 1)
        # Cambiar el archivo invalida el ZIP anterior antes de otra generación.
        with patch('streamlit.file_uploader', return_value=BytesIO(libro([fila(CANAL='CREDICARDPOS')]))):
            at.run()
            self.assertFalse(at.get('download_button'))
            at.button[0].click().run()
            self.assertFalse(at.exception)
            self.assertEqual(at.metric[0].value, '0')
        with patch('streamlit.file_uploader', return_value=BytesIO(b'no es un XLSX')):
            at.run()
            at.button[0].click().run()
            self.assertEqual(len(at.error), 1)
            self.assertFalse(at.get('download_button'))

    def test_celdas_ausentes_del_xml_y_dimensiones_sobredimensionadas(self):
        w = load_workbook(BytesIO(libro([fila()])))
        s = w['VENTAS']
        for col in (5, 9, 16):
            s._cells.pop((2, col), None)
        s.cell(2, 2054).number_format = '0.00'
        b = BytesIO()
        w.save(b)
        _, libros = salidas(b.getvalue())
        s = salida(libros, 'REGION ORIENTE')
        self.assertLess(s.max_column, len(HEADERS))
        self.assertIsNone(valores(s)[0]['OBSERVACION'])

    def test_viernes_toda_la_semana_y_cambio_ano(self):
        for dia in range(14, 21):
            self.assertEqual(p.viernes_misma_semana(date(2026, 9, dia)), date(2026, 9, 18))
        self.assertEqual(p.viernes_misma_semana(date(2026, 12, 31)), date(2027, 1, 1))

    def test_aliases_y_once_agentes_sin_filtro_de_importe(self):
        aliases = ['CENTRO', 'ORIENTE', 'OCCIDENTE', 'CENTRO TIPO II', 'CREDICARDPOSGRANPRO',
                   'MULTITIENDA', 'INVERSIONES TPOS', 'POSMARACAY26', 'POSMGTA25 CA', 'VENEPOS', 'VIRTUALNET']
        r, libros = salidas(libro([fila(CANAL=a, **{'MONTO COMISION AGENTE AUTORIZADO $': None}) for a in aliases]))
        self.assertEqual(r.archivos_generados, 11)
        for a in p.AGENTES:
            self.assertEqual(len(valores(salida(libros, a))), 1)
            self.assertIsNone(valores(salida(libros, a))[0]['MONTO COMISION AGENTE AUTORIZADO $'])

    def test_sin_filtro_fecha_estatus_ni_historicas(self):
        filas = [fila(FECHA=date(2026, mes, 1), ESTATUS=e) for mes, e in [(1, 'PAGADO'), (8, 'DESINSTALADO'), (9, 'Pendiente')]]
        _, libros = salidas(libro(filas))
        v = valores(salida(libros, 'REGION ORIENTE'))
        self.assertEqual(len(v), 3)
        self.assertEqual([r['ESTATUS'] for r in v], ['PAGADO', 'DESINSTALADO', 'Pendiente'])
        self.assertTrue(all(r['MONTO TX AGOSTO'] == 1501 for r in v))

    def test_regionales_independientes_del_canal(self):
        filas = [fila(CANAL=c, **{'VENDEDOR / FREELANCE': p.REGIONALES[0], 'MONTO COMISION VENDEDOR/FREELANCE $': 10}) for c in ('CREDICARDPOS', 'REGION ORIENTE')]
        r, libros = salidas(libro(filas))
        self.assertEqual(len(valores(salida(libros, p.REGIONALES[0]))), 2)
        self.assertEqual(r.archivos_generados, 2)

    def test_barras_exactas_y_ambiguas(self):
        filas = [fila(CANAL='CREDICARDPOS', **{'VENDEDOR / FREELANCE': nombre, 'MONTO COMISION VENDEDOR/FREELANCE $': 10}) for nombre in
                 [p.REGIONALES[0] + '/Persona Ficticia', p.REGIONALES[0] + '/' + p.REGIONALES[1], 'Prefijo ' + p.REGIONALES[0]]]
        r, libros = salidas(libro(filas))
        self.assertEqual(len(valores(salida(libros, p.REGIONALES[0]))), 1)
        self.assertEqual(r.advertencias['Filas con comisión regional sin destinatario inequívoco'], 2)
        self.assertEqual(r.archivos_generados, 1)
        self.assertEqual(valores(salida(libros, p.REGIONALES[0]))[0]['VENDEDOR / FREELANCE'],
                         p.REGIONALES[0] + '/Persona Ficticia')

    def test_freelance_con_barra_conserva_texto_original_completo(self):
        textos = ['PERSONA FICTICIA / Liseth Villanueva',
                  '  liseth   villanueva / PERSONA FICTICIA  ',
                  'PERSONA FICTICIA/Liseth Villanueva/OTRA PERSONA FICTICIA']
        filas = [fila(CANAL='CREDICARDPOS', **{'VENDEDOR / FREELANCE': texto,
                  'MONTO COMISION VENDEDOR/FREELANCE $': 10}) for texto in textos]
        resultado, libros = salidas(libro(filas))
        exportadas = valores(salida(libros, 'Liseth Villanueva'))
        self.assertEqual([r['VENDEDOR / FREELANCE'] for r in exportadas], textos)
        self.assertEqual([r['MONTO COMISION VENDEDOR/FREELANCE $'] for r in exportadas], [10, 10, 10])
        self.assertEqual(resultado.archivos_generados, 1)
        self.assertFalse(resultado.advertencias)

    def test_bancos_por_vendedor_y_monto_no_banco_general(self):
        filas = [fila(CANAL='CREDICARDPOS', BANCO='BANCARIBE', **{'VENDEDOR BANCO': b, 'MONTO COMISION BANCO $': m})
                 for b, m in [('TESORO', 10), ('BANCO DEL TESORO', 7.5), ('BANCARIBE', 10), ('', None), ('BANCARIBE', 0)]]
        r, libros = salidas(libro(filas))
        self.assertEqual(r.archivos_generados, 2)
        self.assertEqual(len(valores(salida(libros, 'TESORO'))), 2)
        self.assertEqual(len(valores(salida(libros, 'BANCARIBE'))), 1)

    def test_misma_venta_en_tres_grupos_y_totales_por_destinatario(self):
        row = fila(BANCO='TESORO', **{p.JORNADA: 'JORNADA DEL TESORO', 'VENDEDOR BANCO': 'TESORO',
                    'MONTO COMISION BANCO $': 10, 'VENDEDOR / FREELANCE': p.REGIONALES[0],
                    'MONTO COMISION VENDEDOR/FREELANCE $': 7.5})
        r, libros = salidas(libro([row]))
        self.assertEqual(r.archivos_generados, 3)
        for nombre, esperado in [('TESORO', 10), (p.REGIONALES[0], 7.5), ('REGION ORIENTE', 20)]:
            s = salida(libros, nombre)
            for cells in s:
                for c in cells:
                    if c.value == 'US$':
                        self.assertEqual(s.cell(c.row, c.column + 1).value, esperado)
                        # Texto blanco de encabezado conserva también su fondo oscuro.
                        self.assertEqual(c.fill.fgColor.rgb, '00123456')
                    if c.value in ('TASA', 'BS.'):
                        self.assertIsNone(s.cell(c.row, c.column + 1).value)

    def test_access_nunca_numero_y_excepciones_compartidas(self):
        filas = [fila(), fila(**{'REGISTRO DE OPERADORES': None}),
                 fila(EQUIPO='Pinpagos', **{'REGISTRO DE OPERADORES': None}),
                 fila(BANCO='TESORO', **{p.JORNADA: 'JORNADA DEL TESORO'}),
                 fila(BANCO='BANCARIBE', VENDEDOR='Persona Ficticia/Otra Ficticia'),
                 fila(**{'REGISTRO DE OPERADORES': 'SI'}),
                 fila(**{'REGISTRO DE OPERADORES': 0}),
                 fila(**{'REGISTRO DE OPERADORES': 'NO', 'REGISTROS DE OPERADORES ACCESS': 111111})]
        _, libros = salidas(libro(filas))
        v = valores(salida(libros, 'REGION ORIENTE'))
        self.assertEqual([r['ACCESS COMMERCE'] for r in v], ['SI', 'NO', 'NO APLICA', 'NO APLICA', 'NO APLICA', 'SI', 'NO', 'SI'])
        self.assertTrue(all(not any('REGISTRO' in str(k) for k in r) for r in v))

    def test_limpieza_comun_y_especifica_tres_grupos(self):
        row = fila(**{'VENDEDOR BANCO': 'TESORO', 'MONTO COMISION BANCO $': 10,
                      'VENDEDOR / FREELANCE': p.REGIONALES[0], 'MONTO COMISION VENDEDOR/FREELANCE $': 10})
        _, libros = salidas(libro([row]))
        for grupo, nombre in [('AGENTES AUTORIZADOS', 'REGION ORIENTE'), ('REGIONALES', p.REGIONALES[0]), ('BANCOS', 'TESORO')]:
            s = salida(libros, nombre)
            h = [c.value for c in s[1]]
            self.assertEqual(h.count('FECHA DE PAGO'), 1)
            for prohibido in ['SERIAL', 'SERIAL.1', 'CONCATENAR', 'CONCATENAR.1', 'FECHA DE ARCHIVO', 'NUMERO DE CUENTA',
                               'GARANTIA DE POS', 'TOTAL VENTA ESTIMADO $$', 'MONTO TOTAL A PAGAR $', '__ROW_ID']:
                self.assertNotIn(prohibido, h)
            for g, par in p.PARES.items():
                for c in par:
                    self.assertEqual(c in h, g == grupo)
            self.assertEqual([HEADERS.index(c) for c in h if c != 'ACCESS COMMERCE'], sorted(HEADERS.index(c) for c in h if c != 'ACCESS COMMERCE'))

    def test_formato_celdas_dimensiones_y_filtro(self):
        datos = libro([fila(), fila()])
        original = load_workbook(BytesIO(datos))['VENTAS']
        _, libros = salidas(datos)
        s = salida(libros, 'REGION ORIENTE')
        h = [c.value for c in s[1]]
        for nombre in ['ESTATUS', 'FECHA', 'MONTO TX AGOSTO']:
            a, b = HEADERS.index(nombre) + 1, h.index(nombre) + 1
            for r in (1, 2, 3):
                for atributo in ('font', 'fill', 'border', 'alignment', 'number_format'):
                    self.assertEqual(copy(getattr(original.cell(r, a), atributo)), copy(getattr(s.cell(r, b), atributo)))
            self.assertEqual(original.column_dimensions[original.cell(1, a).column_letter].width, s.column_dimensions[s.cell(1, b).column_letter].width)
        self.assertEqual([s.row_dimensions[r].height for r in (1, 2, 3)], [38, 27, 28])
        self.assertTrue(s.auto_filter.ref.endswith('3'))

    def test_zip_plano_una_hoja_18_beneficiarios(self):
        filas = [fila(CANAL=a) for a in p.AGENTES]
        filas += [fila(CANAL='CREDICARDPOS', **{'VENDEDOR / FREELANCE': n, 'MONTO COMISION VENDEDOR/FREELANCE $': 10}) for n in p.REGIONALES]
        filas += [fila(CANAL='CREDICARDPOS', **{'VENDEDOR BANCO': b, 'MONTO COMISION BANCO $': 10}) for b in p.BANCOS]
        r, libros = salidas(libro(filas))
        self.assertEqual(r.archivos_generados, 18)
        self.assertEqual(r.nombre_zip, 'Pagos de comisiones 18-09-2026.zip')
        for nombre, w in libros.items():
            self.assertNotIn('/', nombre)
            self.assertEqual(w.sheetnames, ['VENTAS'])

    def test_ceros_resumen_sin_excel_vacio(self):
        r, libros = salidas(libro([fila(CANAL='CREDICARDPOS')]))
        self.assertEqual(len(r.resumen), 18)
        self.assertEqual(r.archivos_generados, 0)
        self.assertFalse(libros)

    def test_validacion_hoja_y_columnas_faltantes_ambiguas(self):
        for datos in [libro([fila()], hoja='OTRA'), libro([fila()], headers=[h for h in HEADERS if h != 'VENDEDOR BANCO']),
                      libro([fila()], headers=HEADERS + ['CANAL.1']),
                      libro([fila()], headers=[h for h in HEADERS if h not in p.ACCESS])]:
            with self.assertRaises(p.ErrorSeparador):
                p.separar_pagos(datos)

    def test_bloque_incompleto_se_bloquea(self):
        with self.assertRaises(p.ErrorSeparador):
            p.separar_pagos(libro([fila()], headers=[h for h in HEADERS if h != 'TOTAL VENTA ESTIMADO $$']))

    def test_variantes_encabezados_y_columnas_movidas(self):
        h = HEADERS.copy()
        h.remove('CANAL')
        h.insert(0, ' canal ')
        row = fila(**{' canal ': 'ORIENTE'})
        _, libros = salidas(libro([row], headers=h))
        self.assertEqual(len(valores(salida(libros, 'REGION ORIENTE'))), 1)

    def test_formula_sin_cache_no_exporta_vacio(self):
        with self.assertRaisesRegex(p.ErrorSeparador, 'fórmulas sin resultado'):
            p.separar_pagos(libro([fila(**{'MONTO COMISION AGENTE AUTORIZADO $': '=10+10'})]))

    def test_importes_invalidos_se_bloquean_sin_datos_sensibles(self):
        for valor in ['importe ficticio inválido', -10, 'Infinity']:
            with self.assertRaises(p.ErrorSeparador) as error:
                p.separar_pagos(libro([fila(**{'MONTO COMISION AGENTE AUTORIZADO $': valor})]))
            self.assertNotIn('ficticio', str(error.exception))

    def test_input_intacto_y_sin_llamar_motor(self):
        datos = libro([fila()])
        copia = bytes(datos)
        with patch('reglas_comisiones.aplicar_motor_comisiones', side_effect=AssertionError('No recalcular')):
            r = p.separar_pagos(datos)
        self.assertEqual(datos, copia)
        self.assertEqual(r.archivos_generados, 1)

    def test_texto_con_igual_no_se_convierte_en_formula(self):
        datos = libro([fila(OBSERVACION='texto')])
        w = load_workbook(BytesIO(datos))
        c = w['VENTAS'].cell(2, HEADERS.index('OBSERVACION') + 1)
        c.value = '=texto ficticio'
        c.data_type = 's'
        b = BytesIO()
        w.save(b)
        _, libros = salidas(b.getvalue())
        s = salida(libros, 'REGION ORIENTE')
        col = [c.value for c in s[1]].index('OBSERVACION') + 1
        self.assertEqual(s.cell(2, col).data_type, 's')


if __name__ == '__main__':
    unittest.main()

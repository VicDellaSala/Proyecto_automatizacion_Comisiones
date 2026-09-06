import io
import unittest

import pandas as pd
from openpyxl import Workbook, load_workbook

import procesamiento as p


# Encabezados del bloque operativo; todos los valores de prueba son ficticios.
ENCABEZADOS = [
    "SERIAL", "CONCATENAR", "fecha de pago", "AFILIADO", "ESTATUS",
    "OBSERVACION", "MES DE CIERRE", "FECHA DE SOLICITUD RECIBIDA",
    "FECHA DE ARCHIVO", "CANAL", "CONCATENAR", "AFILIADO", "TERMINAL",
    "FECHA", "VENDEDOR", "EQUIPO", "SERIAL", "CON TX",
    "REGISTROS DE OPERADORES ACCESS",
]


def maestro_ficticio():
    wb = Workbook()
    ws = wb.active
    ws.title = "VENTAS"
    ws.append(ENCABEZADOS)
    ws.append(["AUX", "1119", None, "111", "Pendiente", "Nota ficticia",
               None, None, None, None, "2221", "222", 1, None, None,
               "Castle Dynamo", "PRINCIPAL", "", "SI"])
    wb.create_sheet("OTRA").append(["=1+1"])
    archivo = io.BytesIO()
    wb.save(archivo)
    return archivo


class EstructuraComisionesTests(unittest.TestCase):
    def test_identidad_operativa_en_preparacion_y_recalculo(self):
        base = p.preparar_comisiones(maestro_ficticio())
        self.assertEqual(base.loc[0, "__AFILIADO"], "222")
        self.assertEqual(base.loc[0, "__CONCATENAR"], "2221")
        self.assertEqual(base.loc[0, "__SERIAL_COMISION"], "PRINCIPAL")
        r34 = pd.DataFrame({"__CONCATENAR": ["2221"], "__MONTO_TX": [1],
                            "__SERIAL_R34": ["PRINCIPAL"]})
        final = p.recalcular_comisiones(base, r34, {"111"})
        self.assertEqual(final.loc[0, "__ACCESS_CALCULADO"], "NO")
        self.assertEqual(final.loc[0, "__ESTADO_TX_CALCULADO"], "C/P SIN TX")
        self.assertEqual(final.loc[0, "__MOTIVO_REVISION"], "")
        final = p.recalcular_comisiones(base, r34, {"222"})
        self.assertEqual(final.loc[0, "__ACCESS_CALCULADO"], "SI")
        # Las correcciones del bloque principal prevalecen sobre auxiliares e internos viejos.
        final.loc[0, "AFILIADO.1"] = "333"
        final.loc[0, "TERMINAL"] = 2
        final.loc[0, "SERIAL.1"] = "CORREGIDO"
        final = p.recalcular_comisiones(final, pd.DataFrame(), {"222"})
        self.assertEqual(final.loc[0, "__AFILIADO"], "333")
        self.assertEqual(final.loc[0, "__CONCATENAR"], "3332")
        self.assertEqual(final.loc[0, "__SERIAL_COMISION"], "CORREGIDO")
        self.assertEqual(final.loc[0, "__ACCESS_CALCULADO"], "NO")
        self.assertEqual(final.loc[0, "AFILIADO"], "111")

    def test_no_elije_ultima_columna_sin_bloque_conocido(self):
        with self.assertRaises(ValueError):
            p.recalcular_comisiones(pd.DataFrame({"AFILIADO": ["111"],
                "AFILIADO.1": ["222"], "TERMINAL": [1]}), pd.DataFrame(), {"222"})

    def test_bloque_se_resuelve_aunque_cambie_posicion(self):
        archivo = maestro_ficticio()
        wb = load_workbook(archivo)
        wb['VENTAS'].insert_cols(1)
        wb['VENTAS']['A1'] = 'OTRO CAMPO'
        nuevo = io.BytesIO()
        wb.save(nuevo)
        base = p.preparar_comisiones(nuevo)
        self.assertEqual(base.loc[0, "__AFILIADO"], "222")

    def test_integracion_exporta_ambos_bloques_y_conserva_encabezados(self):
        archivo = maestro_ficticio()
        base = p.preparar_comisiones(archivo)
        ventas = pd.DataFrame({"AFILIADO": ["222", "444"], "TERMINAL": [1, 2],
            "SERIAL": ["PRINCIPAL", "NUEVO"], "EQUIPO": ["Castle Dynamo"] * 2,
            "CONCATENAR": ["NO USAR", "NO USAR"],
            "__AFILIADO": ["222", "444"], "__TERMINAL": ["1", "2"],
            "__CONCATENAR": ["2221", "4442"], "__EQUIPO_STD": ["Castle Dynamo"] * 2})
        combinado, nuevas, existentes = p.integrar_ventas(base, ventas)
        self.assertEqual((len(nuevas), len(existentes)), (1, 1))
        final = p.recalcular_comisiones(combinado, pd.DataFrame(), {"444"})
        for col in ['AFILIADO', 'AFILIADO.1']:
            self.assertEqual(final.loc[1, col], '444')
        for col in ['CONCATENAR', 'CONCATENAR.1']:
            self.assertEqual(final.loc[1, col], '4442')
        self.assertEqual(final.loc[1, '__ACCESS_CALCULADO'], 'SI')
        salida = p.generar_excel_resultado({'bytes_comisiones_original': archivo.getvalue(),
            'hoja_comisiones': 'VENTAS', 'cantidad_original': 1, 'final': final})
        wb = load_workbook(io.BytesIO(salida))
        self.assertEqual([c.value for c in wb['VENTAS'][1]], ENCABEZADOS)
        self.assertEqual(wb['VENTAS']['L3'].value, '444')
        self.assertEqual(wb['VENTAS']['K3'].value, '4442')
        self.assertEqual(wb['OTRA']['A1'].value, '=1+1')

    def test_modalidad_reporte_alimenta_cxc_solo_en_nuevas(self):
        base = p.preparar_comisiones(maestro_ficticio())
        base['ESTATUS CXC'] = 'COMODATO'
        ventas = pd.DataFrame({'AFILIADO': ['444'], 'TERMINAL': [2],
            'EQUIPO': ['Castle Dynamo'], '__AFILIADO': ['444'], '__TERMINAL': ['2'],
            '__CONCATENAR': ['4442'], '__EQUIPO_STD': ['Castle Dynamo'],
            'DECONTADO / FINANCIAMIENTO': ['al contado']})
        nuevas = p.crear_filas_nuevas(ventas, base)
        self.assertEqual(nuevas.loc[0, 'ESTATUS CXC'], 'AL CONTADO')
        self.assertEqual(base.loc[0, 'ESTATUS CXC'], 'COMODATO')

    def test_sufijo_tx_se_conserva_sin_inferir_periodo(self):
        base = p.preparar_comisiones(maestro_ficticio())
        columnas = ['Monto_Trans_Acum_bs_mes', 'Monto_Trans_Acum_bs_mes_1']
        for col in columnas:
            base[col] = None
        ventas = pd.DataFrame({
            'AFILIADO': ['444'], 'TERMINAL': [2], 'EQUIPO': ['Pinpagos'],
            '__AFILIADO': ['444'], '__TERMINAL': ['2'], '__CONCATENAR': ['4442'],
            '__EQUIPO_STD': ['Pinpagos'], columnas[0]: [17], columnas[1]: [29],
        })
        nuevas = p.crear_filas_nuevas(ventas, base)
        self.assertEqual(nuevas.loc[0, columnas[0]], 17)
        self.assertEqual(nuevas.loc[0, columnas[1]], 29)
        # Tener solo la columna sin sufijo no autoriza rellenar la de otro período.
        nuevas = p.crear_filas_nuevas(ventas.drop(columns=[columnas[1]]), base)
        self.assertTrue(pd.isna(nuevas.loc[0, columnas[1]]))


class PertenenciaTests(unittest.TestCase):
    def test_lista_blanca_por_chunks(self):
        permitidas = ['CREDICARD POS', 'CREDICARDPOS', ' credicardpos   cdm ']
        excluidas = ['CREDICARD CENTRO', 'CREDICARD ORIENTE', 'CREDICARD OCCIDENTE',
                     'CREDICARDPOS OCCIDENT CDM', 'OTRO CREDICARDPOS', 'CREDICARDPOS CDM EXTRA', '']
        texto = 'PERTENENCIA;AFIPOS;MONTO_TRANS_BS_ACUM_MES\n' + ''.join(
            f'{valor};{i};1\n' for i, valor in enumerate(permitidas + excluidas, 1))
        resultado, detalle = p.procesar_csv_r34(io.BytesIO(texto.encode()), 'ficticio.csv', chunksize=2)
        self.assertEqual(resultado['__CONCATENAR'].tolist(), ['1', '2', '3'])
        self.assertEqual(detalle['filas_leidas'], 10)
        self.assertEqual(detalle['filas_credicardpos'], 3)


if __name__ == '__main__':
    unittest.main()

import io
import unittest
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import procesamiento as p
from tx_mensual import numero_tx, preparar_periodos_ventas, unir_historial_ventas, aplicar_historial
from test_estructura_comisiones import maestro_ficticio


class TxMensualTests(unittest.TestCase):
    def base(self, mes=9):
        b=p.preparar_comisiones(maestro_ficticio())
        b['__ORIGEN']='VENTAS_NUEVAS'; b['__ANO_REPORTE']=2026; b['__MES_REPORTE']=mes
        b['OBSERVACION']=''
        return b

    def r34(self, monto=2000, mes=9, serial='PRINCIPAL', error=''):
        n,e=numero_tx(monto)
        return pd.DataFrame([{'__CONCATENAR':'2221','__ANO_R34':2026,'__MES_R34':mes,
            '__MONTO_TX':n,'__MOTIVO_MONTO_TX':error or e,'__SERIAL_R34':serial}])

    def test_limites_y_diagnostico(self):
        for valor,estado,motivo in [(0,'SIN TX',''),('0','SIN TX',''),('-','SIN TX',''),(' - ','SIN TX',''),
            (0.01,'C/P SIN TX',''),(1,'C/P SIN TX',''),(999.99,'C/P SIN TX',''),(1000,'REVISAR 1000',''),
            (1000.01,'CON_TX',''),('', 'N/A','MONTO_TX_VACIO'),(None,'N/A','MONTO_TX_VACIO'),
            (float('nan'),'N/A','MONTO_TX_VACIO'),('roto12','N/A','MONTO_TX_INVALIDO')]:
            self.assertEqual(numero_tx(valor)[1],motivo)
            self.assertEqual(p.estado_transaccion(valor),estado)

    def test_calendario_y_cambio_ano(self):
        for ano,mes,prev in [(2026,8,(2026,7)),(2026,9,(2026,8)),(2026,10,(2026,9)),(2027,1,(2026,12))]:
            d=pd.DataFrame({'FECHA REPORTE':[pd.Timestamp(ano,mes,1)],
                'Monto_Trans_Acum_bs_mes':[12],'Monto_Trans_Acum_bs_mes_1':[8]})
            d=preparar_periodos_ventas(d,'ficticio.xlsx')
            self.assertEqual(d.iloc[0]['__TX_VENTAS'],{(ano,mes):(12,''),prev:(8,'')})
            self.assertEqual(d.iloc[0]['__MES_REPORTE'],mes)

    def test_lookup_no_elige_mayor_entre_meses(self):
        r=pd.concat([self.r34(500,8),self.r34(1500,9),self.r34(600,8)],ignore_index=True)
        lookup=p.crear_lookup_r34(r)
        self.assertEqual(lookup[('2221',2026,8)]['__MONTO_TX'],600)
        for mes,estado,monto in [(8,'C/P SIN TX',600),(9,'CON_TX',1500),(10,'N/A',None)]:
            out=p.recalcular_comisiones(self.base(mes),r,set())
            self.assertEqual(out.loc[0,'CON TX'],'CON_TX')
            self.assertEqual(out.loc[0,'MONTO TX AGOSTO'],600)
            self.assertEqual(out.loc[0,'MONTO TX SEPTIEMBRE'],1500)
            if monto is None: self.assertIn('NO_ENCONTRADO_R34',out.loc[0,'__MOTIVO_TX'])
            else:self.assertEqual(out.loc[0,'__MONTO_TX_R34'],monto)

    def test_r34_prioritario_historia_discrepancia(self):
        b=self.base();b['__TX_VENTAS']=[{(2026,8):(500,''),(2026,9):(900,'')}]
        out=p.recalcular_comisiones(b,self.r34(1500),set())
        self.assertEqual(out.loc[0,'MONTO TX AGOSTO'],500)
        self.assertEqual(out.loc[0,'MONTO TX SEPTIEMBRE'],1500)
        self.assertEqual(out.loc[0,'__DIFERENCIA_TX_FUENTES'][(2026,9)]['diferencia'],600)
        self.assertEqual(out.loc[0,'CON TX'],'CON_TX')
        sin=p.recalcular_comisiones(b,self.r34(2000,8),set())
        self.assertEqual(sin.loc[0,'CON TX'],'CON_TX')
        self.assertEqual(sin.loc[0,'MONTO TX SEPTIEMBRE'],900)

    def test_observaciones_na_y_serial(self):
        for r,motivo,texto in [(pd.DataFrame(),'NO_ENCONTRADO_R34','NO ENCONTRADO EN R34'),
            (self.r34(None),'MONTO_TX_VACIO','MONTO TX VACIO'),
            (self.r34('roto'),'MONTO_TX_INVALIDO','MONTO TX INVALIDO'),
            (self.r34(None,serial=''),'SERIAL_SIN_FUENTE_CONFIABLE','SERIAL SIN FUENTE CONFIABLE')]:
            b=self.base();o=p.recalcular_comisiones(b,r,set())
            self.assertEqual(o.loc[0,'CON TX'],'N/A')
            self.assertIn(motivo,o.loc[0,'__MOTIVO_TX'])
            self.assertEqual(o.loc[0,'OBSERVACION'],'N/A - '+texto)
            for origen in ['COMISIONES','VENTAS_NUEVAS']:
                b['__ORIGEN']=origen;b['OBSERVACION']='Nota ficticia conservada'
                self.assertEqual(p.recalcular_comisiones(b,r,set()).loc[0,'OBSERVACION'],'Nota ficticia conservada')
        for serial,motivo in [('DISTINTO','SERIAL_DIFERENTE_R34'),('','SERIAL_SIN_FUENTE_CONFIABLE')]:
            o=p.recalcular_comisiones(self.base(),self.r34(2000,serial=serial),set())
            self.assertEqual(o.loc[0,'CON TX'],'CON_TX')
            self.assertIn(motivo,o.loc[0,'__MOTIVO_TX'])
            self.assertFalse(o.loc[0,'OBSERVACION'].startswith('N/A -'))

    def test_historial_sobrevive_deduplicacion(self):
        partes=[]
        for mes in [8,9,10]:
            d=pd.DataFrame({'__CONCATENAR':['2221'],'FECHA REPORTE':[pd.Timestamp(2026,mes,1)],
                'Monto_Trans_Acum_bs_mes':[mes]})
            partes.append(preparar_periodos_ventas(d,'ficticio.xlsx'))
        d=pd.concat(partes,ignore_index=True);unir_historial_ventas(d)
        self.assertEqual(set(d.iloc[0]['__TX_VENTAS']),{(2026,8),(2026,9),(2026,10)})

    def test_columnas_exportacion_estilos_sin_duplicados(self):
        f=maestro_ficticio();w=load_workbook(f);s=w['VENTAS']
        c=s.max_column+1;s.cell(1,c,'MONTO TX AGOSTO');s.cell(2,c,7)
        s.cell(1,c).fill=PatternFill('solid',fgColor='FF00FF00')
        s.cell(2,c).number_format='#,##0.00';s.column_dimensions[s.cell(1,c).column_letter].width=23
        f=io.BytesIO();w.save(f);original=f.getvalue()
        b=p.preparar_comisiones(f);b['__MES_REPORTE']=9;b['__ANO_REPORTE']=2026;b['__ORIGEN']='VENTAS_NUEVAS'
        o=p.recalcular_comisiones(b,self.r34(1500),set())
        o=p.recalcular_comisiones(o,self.r34(1500),set())
        self.assertEqual(list(o).count('MONTO TX SEPTIEMBRE'),1)
        result={'final':o,'cantidad_original':1,'bytes_comisiones_original':original,'hoja_comisiones':'VENTAS'}
        export=p.generar_excel_resultado(result);w=load_workbook(io.BytesIO(export));s=w['VENTAS']
        self.assertEqual(s.cell(1,c+1).value,'MONTO TX SEPTIEMBRE')
        self.assertEqual(s.cell(2,c+1).value,1500)
        self.assertEqual(s.cell(2,c+1).number_format,s.cell(2,c).number_format)
        self.assertEqual(s.cell(1,c+1).fill.fgColor.rgb,s.cell(1,c).fill.fgColor.rgb)
        self.assertEqual(s.column_dimensions[s.cell(1,c+1).column_letter].width,23)
        self.assertEqual(w['OTRA']['A1'].value,'=1+1')
        self.assertFalse(any(str(cell.value).startswith('__') for cell in s[1]))

    def test_acumulativo_no_retrocede_y_no_na_posterior(self):
        for anterior,actual,esperado in [(1500,None,'CON_TX'),(1500,0,'CON_TX'),(1500,'-','CON_TX'),
            (500,None,'C/P SIN TX'),(1000,None,'REVISAR 1000'),(0,None,'SIN TX'),(None,None,'N/A')]:
            b=self.base();b['MONTO TX AGOSTO']=anterior;b['MONTO TX JULIO']=None if anterior is None else '-'
            o=p.recalcular_comisiones(b,self.r34(actual),set())
            self.assertEqual(o.loc[0,'CON TX'],esperado)
            self.assertEqual(o.loc[0,'ESTATUS'],b.loc[0,'ESTATUS'])
            if esperado!='N/A':self.assertFalse(o.loc[0,'OBSERVACION'].startswith('N/A -'))
        b=self.base();b['CON TX']='CON_TX'
        self.assertEqual(p.recalcular_comisiones(b,self.r34(0),set()).loc[0,'CON TX'],'CON_TX')

    def test_r34_periodos_validacion_y_respaldo(self):
        from tx_mensual import validar_periodo_r34
        self.assertIsNone(validar_periodo_r34(2026,13))
        self.assertIsNone(validar_periodo_r34(2026,9.5))
        for ano,mes,esperado in [('2026','9',(2026,9)),('2027','1',(2027,1)),('','9',None),('2026','13',None)]:
            csv=f'PERTENENCIA;AFIPOS;SERIAL;MONTO_TRANS_BS_ACUM_MES;ANO_PROCESO;MES_PROCESO\nCREDICARDPOS;2221;PRINCIPAL;1;{ano};{mes}\n'
            df,d=p.procesar_csv_r34(io.BytesIO(csv.encode()),'ficticio.csv',1)
            self.assertEqual(d['requiere_periodo_manual'],esperado is None)
            if esperado:self.assertEqual(d['periodos'],[esperado])
            else:
                df,d=p.procesar_csv_r34(io.BytesIO(csv.encode()),'ficticio.csv',1,(2026,10))
                self.assertFalse(d['requiere_periodo_manual'])
                self.assertEqual(d['periodos'],[(2026,10)])
        csv='PERTENENCIA;AFIPOS;SERIAL;MONTO_TRANS_BS_ACUM_MES;ANO_PROCESO;MES_PROCESO\nCREDICARDPOS;2221;PRINCIPAL;500;2026;8\nCREDICARDPOS;2221;PRINCIPAL;1500;2026;9\n'
        df,d=p.procesar_csv_r34(io.BytesIO(csv.encode()),'ficticio.csv',1)
        self.assertEqual(d['periodos'],[(2026,8),(2026,9)])
        self.assertIsNone(p.determinar_mes_r34([d]))
        self.assertEqual(len(p.crear_lookup_r34(df)),2)

    def test_r34_anterior_actualiza_solo_periodos_explicitos(self):
        for ano,mes,prev in [(2026,8,(2026,7)),(2026,9,(2026,8)),(2026,10,(2026,9)),(2027,1,(2026,12))]:
            csv=f'PERTENENCIA;AFIPOS;SERIAL;Monto_Trans_Acum_bs_mes;Monto_Trans_Acum_bs_mes_1;ANO_PROCESO;MES_PROCESO\nCREDICARDPOS;2221;PRINCIPAL;300;1500;{ano};{mes}\n'
            df,d=p.procesar_csv_r34(io.BytesIO(csv.encode()),'ficticio.csv',1)
            lookup=p.crear_lookup_r34(df)
            self.assertEqual(lookup[('2221',*prev)]['__MONTO_TX'],1500)
            self.assertEqual(lookup[('2221',ano,mes)]['__MONTO_TX'],300)
            if mes==9:
                b=self.base();b['MONTO TX JUNIO']=0;b['MONTO TX JULIO']='-';b['MONTO TX AGOSTO']=500
                o=p.recalcular_comisiones(b,df,set())
                self.assertEqual(o.loc[0,'MONTO TX AGOSTO'],1500)
                self.assertEqual(o.loc[0,'MONTO TX SEPTIEMBRE'],300)
                self.assertEqual(o.loc[0,'MONTO TX JULIO'],'-')
                self.assertEqual(o.loc[0,'MONTO TX JUNIO'],0)
                self.assertEqual(o.loc[0,'CON TX'],'CON_TX')

    def test_periodo_ventas_no_se_inventa_del_nombre(self):
        d=pd.DataFrame({'FECHA REPORTE':['incorrecta'],'Monto_Trans_Acum_bs_mes':[5]})
        d=preparar_periodos_ventas(d,'SEPTIEMBRE 2026.xlsx')
        self.assertIsNone(d.loc[0,'__MES_REPORTE'])
        self.assertEqual(d.loc[0,'__TX_VENTAS'],{})
        self.assertTrue(d.loc[0,'__DIAGNOSTICO_PERIODO'])

    def test_columna_insertada_referencias_tabla_y_otra_hoja(self):
        from openpyxl.worksheet.table import Table
        w=load_workbook(maestro_ficticio());ws=w['VENTAS'];c=ws.max_column+1
        ws.cell(1,c,'MONTO TX AGOSTO');ws.cell(2,c,500)
        ws.cell(1,c+1,'OTRO CAMPO');ws.cell(2,c+1,9)
        ws.cell(1,c+2,'FORMULA');ws.cell(2,c+2,f'={ws.cell(2,c+1).coordinate}+1')
        w['OTRA']['A1']=f"='VENTAS'!{ws.cell(2,c+1).coordinate}"
        # Tabla pequeña con encabezados únicos alrededor del punto de inserción.
        ws.add_table(Table(displayName='TablaFicticia',ref=f'{ws.cell(1,c).coordinate}:{ws.cell(2,c+2).coordinate}'))
        f=io.BytesIO();w.save(f);b=p.preparar_comisiones(f)
        b['__MES_REPORTE']=9;b['__ANO_REPORTE']=2026;b['__ORIGEN']='VENTAS_NUEVAS'
        out=p.recalcular_comisiones(b,self.r34(300),set())
        result={'final':out,'cantidad_original':1,'bytes_comisiones_original':f.getvalue(),'hoja_comisiones':'VENTAS'}
        w=load_workbook(io.BytesIO(p.generar_excel_resultado(result)));ws=w['VENTAS']
        self.assertEqual(ws.cell(1,c+1).value,'MONTO TX SEPTIEMBRE')
        self.assertEqual(ws.cell(2,c+2).value,9)
        self.assertEqual(ws.cell(2,c+3).value,f'={ws.cell(2,c+2).coordinate}+1')
        self.assertEqual(w['OTRA']['A1'].value,f"='VENTAS'!{ws.cell(2,c+2).coordinate}")
        self.assertEqual(len(ws.tables['TablaFicticia'].tableColumns),4)

    def test_reporte_actualiza_anterior_existente_sin_tocar_julio(self):
        b=self.base();b['MONTO TX JULIO']=77;b['MONTO TX AGOSTO']=500
        b['__TX_VENTAS']=[{(2026,8):(1500,''),(2026,9):(300,'')}]
        o=p.recalcular_comisiones(b,pd.DataFrame(),set())
        self.assertEqual(o.loc[0,'MONTO TX AGOSTO'],1500)
        self.assertEqual(o.loc[0,'MONTO TX JULIO'],77)
        self.assertEqual(o.loc[0,'CON TX'],'CON_TX')

    def test_conflicto_anos_no_silencioso(self):
        b=self.base();b['MES DE CIERRE']='ENERO (2026)';b['MONTO TX ENERO']=1
        b['__TX_VENTAS']=[{(2027,1):(20,'')}]
        with self.assertRaisesRegex(ValueError,'dos años'):
            aplicar_historial(b,{})

    def test_paso2_muestra_evidencia_de_otro_periodo_sin_usar_historial_sintetico(self):
        from test_revision_manual import resultados
        from revision_manual import incidencias
        r=resultados(('A','PRINCIPAL'));b=r['final'];b['__MES_REPORTE']=9;b['__ANO_REPORTE']=2026
        r34=pd.concat([self.r34(1500,8,'DIFERENTE'),self.r34(500,9)],ignore_index=True)
        r34['__CONCATENAR']='1011'
        self.assertFalse(incidencias(b,r34)[2])  # Alternativas no prueban un cambio real.
        r34.loc[r34['__MES_R34'].eq(8),'__ES_HISTORIAL_TX'] = True
        self.assertFalse(incidencias(b,r34)[2])
        r34.loc[r34['__MES_R34'].eq(9),'__SERIAL_R34']='OTRO'
        self.assertEqual(len(incidencias(b,r34)[2]),1)

    def test_respaldo_solo_para_registros_invalidos(self):
        csv='PERTENENCIA;AFIPOS;SERIAL;MONTO_TX;ANO_PROCESO;MES_PROCESO\nCREDICARDPOS;2221;PRINCIPAL;5;2026;8\nCREDICARDPOS;3331;OTRO;6;;\n'
        r,d=p.procesar_csv_r34(io.BytesIO(csv.encode()),'ficticio.csv',1,(2026,9))
        self.assertEqual(d['periodos'],[(2026,8),(2026,9)])
        self.assertEqual(r['__MES_R34'].tolist(),[8,9])
        self.assertFalse(d['requiere_periodo_manual'])

    def test_pipeline_pide_respaldo_sin_reconstruir_maestro(self):
        from unittest.mock import patch
        from tx_mensual import PeriodoR34Requerido
        detalle=[{'archivo':'ficticio.csv','requiere_periodo_manual':True}]
        with patch.object(p,'procesar_r34',return_value=(pd.DataFrame(),detalle)):
            with self.assertRaises(PeriodoR34Requerido) as error:
                p.procesar_todo([],[],maestro_ficticio(),None)
            self.assertEqual(error.exception.archivos,['ficticio.csv'])

    def test_cadena_calculo_y_columnas_completas(self):
        import xml.etree.ElementTree as ET
        from tx_exportacion import cadena_calculo_xml, formula
        df=pd.DataFrame(columns=['A','MONTO TX AGOSTO','MONTO TX SEPTIEMBRE','B'])
        df.attrs['columnas_antes_tx']=['A','MONTO TX AGOSTO','B']
        xml=b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="C2" i="1"/><c r="C3"/><c r="C2" i="2"/></calcChain>'
        root=ET.fromstring(cadena_calculo_xml(xml,df,'1'))
        self.assertEqual([c.get('r') for c in root],['D2','D3','C2'])
        self.assertEqual(formula('SUM(C:C)',{1:1,2:2,3:4},True),'SUM(D:D)')

    def test_orden_descendente_y_segunda_carga_sin_duplicar(self):
        b=self.base();b['MONTO TX AGOSTO']=1;b['MONTO TX JULIO']=2;b['OTRA COLUMNA']='conservar'
        out=p.recalcular_comisiones(b,self.r34(300),set())
        cols=list(out)
        self.assertEqual(cols[cols.index('MONTO TX SEPTIEMBRE')+1],'MONTO TX AGOSTO')
        out['__MES_REPORTE']=10
        out=p.recalcular_comisiones(out,self.r34(400,10),set())
        cols=list(out)
        self.assertEqual(cols[cols.index('MONTO TX OCTUBRE')+1],'MONTO TX SEPTIEMBRE')
        self.assertEqual(cols.count('MONTO TX SEPTIEMBRE'),1)
        self.assertEqual(out.loc[0,'OTRA COLUMNA'],'conservar')

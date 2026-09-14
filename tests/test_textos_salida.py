import io
import unittest
import zipfile
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

import procesamiento as p
from reglas_comisiones import estandarizar_equipo, normalizar_canal, aplicar_motor_comisiones
from test_revision_manual import resultados


class TextosSalidaTests(unittest.TestCase):
    def test_aliases_en_historicas_y_nuevas_sin_cambiar_otros_campos(self):
        equipos = [('CASTLE','Castle Dynamo'),('CASTTLE','Castle Dynamo'),('DYNAMO','Castle Dynamo'),
                   ('ZAPPY','Zappy S1MINI2'),('SAPPY','Zappy S1MINI2'),('S1MINI2','Zappy S1MINI2'),
                   ('S1 MINI 2','Zappy S1MINI2'),('PINPAGO','Pinpagos'),('PIN PAGO','Pinpagos'),('EQUIPO K','Pinpagos')]
        canales = [('CREDICARDPOSGRANPRO','GRANPRO'),('GRANPRO','GRANPRO'),('POSMGTA25','POSMGTA'),
                   ('POSMGTA25 CA','POSMGTA'),('POSMGTA','POSMGTA'),('OCCIDENTE','REGION OCCIDENTE'),
                   ('ORIENTE','REGION ORIENTE'),('CENTRO','REGION CENTRO'),('CENTRO TIPO II','CENTRO TIPO II'),
                   ('INVERSIONES TPOS','INV TPOS'),('INV TPOS','INV TPOS'),('POSMARACAY26 CA','POSMARACAY'),
                   ('REGION CENTRO','REGION CENTRO'),('REGION ORIENTE','REGION ORIENTE'),
                   ('REGION OCCIDENTE','REGION OCCIDENTE'),('MULTITIENDA','MULTITIENDA'),
                   ('VENEPOS','VENEPOS'),('VIRTUALNET','VIRTUALNET'),('CREDICARDPOS','CREDICARDPOS')]
        for origen in ['COMISIONES','VENTAS_NUEVAS']:
            for (equipo, esperado), (canal, canal_esperado) in [(e,c) for e in equipos for c in canales]:
                d = pd.DataFrame([{'EQUIPO':equipo,'CANAL':canal,'VENDEDOR AGENTE AUTORIZADO':canal,
                    'ESTATUS':'PAGADO','OBSERVACION':'OBSERVACION HISTORICA','MONTO TOTAL A PAGAR $':38.4,
                    'CON TX':'CON_TX','SERIAL':'123456789012','REGISTRO DE OPERADORES':'SI',
                    'MONTO TX AGOSTO':1500,'VENDEDOR / FREELANCE':'Nombre ficticio / Otro ficticio',
                    'VENDEDOR BANCO':'BANCO FICTICIO','CANAL DE VENTA (JORNADA QUE PERTENECE)':'JORNADA DEL TESORO',
                    '__ORIGEN':origen,'__ROW_ID':1}])
                antes=d.copy(deep=True)
                out=p.estandarizar_textos_salida(d)
                self.assertEqual(out.at[0,'EQUIPO'],esperado)
                self.assertEqual(out.at[0,'CANAL'],canal_esperado)
                self.assertEqual(out.at[0,'VENDEDOR AGENTE AUTORIZADO'],canal_esperado)
                otros=[c for c in d if c not in {'EQUIPO','CANAL','VENDEDOR AGENTE AUTORIZADO'}]
                pd.testing.assert_frame_equal(out[otros],antes[otros])
                pd.testing.assert_frame_equal(d,antes)
                pd.testing.assert_frame_equal(p.estandarizar_textos_salida(out),out)

    def test_desconocidos_y_personas_se_conservan_literalmente(self):
        valores=['  Nombre libre á  ','Empresa POSMGTA25 CA','OFICINA - GRAN PRO','',None,pd.NA,float('nan')]
        d=pd.DataFrame({'EQUIPO':valores,'CANAL':valores,'VENDEDOR AGENTE AUTORIZADO':valores})
        pd.testing.assert_frame_equal(p.estandarizar_textos_salida(d),d)
        self.assertEqual(estandarizar_equipo('  Equipo desconocido  ',conservar_desconocidos=True),'  Equipo desconocido  ')
        self.assertEqual(normalizar_canal('  Canal desconocido á  ',conservar_desconocidos=True),'  Canal desconocido á  ')
        self.assertEqual(normalizar_canal('  Canal desconocido á  '),'CANAL DESCONOCIDO A')

    def test_venta_nueva_con_canal_desconocido_no_pierde_texto(self):
        d=resultados()['final'].iloc[[1]].copy()
        d['CANAL']='  Canal desconocido á  '
        out=aplicar_motor_comisiones(d)
        self.assertEqual(out.iloc[0]['CANAL'],'  Canal desconocido á  ')

    def test_xml_exporta_texto_canonico_sin_recalculo_ni_cambio_estructura(self):
        rows=[['CASTLE','CREDICARDPOSGRANPRO','CREDICARDPOSGRANPRO','PAGADO'],
              ['S1 MINI 2','POSMGTA25 CA','Persona ficticia','DESINSTALADO'],
              ['PIN PAGO','OCCIDENTE','OCCIDENTE','Aplica Pago'],
              ['  Equipo libre  ','  Canal libre  ','Persona Granpro','PENDIENTE']]
        d=pd.DataFrame(rows,columns=['EQUIPO','CANAL','VENDEDOR AGENTE AUTORIZADO','ESTATUS'])
        d['OBSERVACION']='OBSERVACION HISTORICA';d['MONTO TOTAL A PAGAR $']=38.4
        d['VENDEDOR / FREELANCE']='CREDICARDPOSGRANPRO';d['VENDEDOR BANCO']='OCCIDENTE'
        d['SERIAL']='123456789012';d['MONTO TX AGOSTO']=1500;d['CON TX']='CON_TX'
        d['CANAL DE VENTA (JORNADA QUE PERTENECE)']='JORNADA DEL TESORO'
        cols=list(d);wb=Workbook();ws=wb.active;ws.title='VENTAS';ws.append(cols)
        for row in d.iloc[:2].itertuples(index=False,name=None):ws.append(row)
        ws['F2']='=19.2*2';ws['A2'].fill=PatternFill('solid',fgColor='FFFF00')
        ws.column_dimensions['A'].width=24;ws.auto_filter.ref='A1:L3'
        wb.create_sheet('OTRA')['A1']='=1+1'
        source=io.BytesIO();wb.save(source)
        d['__ORIGEN']=['COMISIONES']*2+['VENTAS_NUEVAS']*2;d['__ROW_ID']=range(1,5)
        d.attrs['encabezados_comisiones']=dict(zip(cols,cols))
        antes=d.copy(deep=True)
        with patch.object(p,'recalcular_comisiones',side_effect=AssertionError('No recalcular')), \
             patch.object(p,'aplicar_motor_comisiones',side_effect=AssertionError('No ejecutar motor')):
            raw=p.generar_excel_resultado({'final':d,'cantidad_original':2,'hoja_comisiones':'VENTAS',
                                          'bytes_comisiones_original':source.getvalue()})
        pd.testing.assert_frame_equal(d,antes)
        out=load_workbook(io.BytesIO(raw));v=out['VENTAS']
        self.assertEqual([c.value for c in v[1]],cols)
        self.assertEqual([v.cell(i,1).value for i in range(2,6)],['Castle Dynamo','Zappy S1MINI2','Pinpagos','  Equipo libre  '])
        self.assertEqual([v.cell(i,2).value for i in range(2,6)],['GRANPRO','POSMGTA','REGION OCCIDENTE','  Canal libre  '])
        self.assertEqual([v.cell(i,3).value for i in range(2,6)],['GRANPRO','Persona ficticia','REGION OCCIDENTE','Persona Granpro'])
        for i,row in enumerate(rows,2):
            self.assertEqual(v.cell(i,4).value,row[3])
            self.assertEqual(v.cell(i,5).value,'OBSERVACION HISTORICA')
            self.assertEqual(v.cell(i,7).value,'CREDICARDPOSGRANPRO')
            self.assertEqual(v.cell(i,8).value,'OCCIDENTE')
            self.assertEqual(v.cell(i,6).value,'=19.2*2' if i==2 else 38.4)
        self.assertEqual(v['A2'].fill.fgColor.rgb,'00FFFF00')
        self.assertEqual(v.column_dimensions['A'].width,24)
        self.assertEqual(v.auto_filter.ref,'A1:L3')
        with zipfile.ZipFile(source) as original,zipfile.ZipFile(io.BytesIO(raw)) as final:
            for name in original.namelist():
                if name!='xl/worksheets/sheet1.xml':self.assertEqual(original.read(name),final.read(name),name)

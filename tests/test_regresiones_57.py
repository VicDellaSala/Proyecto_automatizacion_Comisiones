import io
import unittest
from datetime import datetime
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

import procesamiento as p
from revision_manual import incidencias, aplicar_decision, preparar_descarga
from test_revision_manual import resultados


class RegresionesTests(unittest.TestCase):
    def test_beneficiario_historico_no_se_reescribe_al_completar_importes(self):
        from reglas_comisiones import aplicar_motor_comisiones
        d = resultados()['final'].iloc[:1].copy()
        d['VENDEDOR AGENTE AUTORIZADO'] = '  Region Centro  '
        for c in d:
            if c.startswith('MONTO COMISION') or c == 'MONTO TOTAL A PAGAR $':
                d[c] = None
        out = aplicar_motor_comisiones(d)
        self.assertEqual(out.at[0,'VENDEDOR AGENTE AUTORIZADO'],'  Region Centro  ')
        self.assertEqual(out.at[0,'MONTO TOTAL A PAGAR $'],20)

    def test_paso1_separa_claves_aun_con_mismo_serial(self):
        r = resultados(('123456789012',)*4, ('COMISIONES','VENTAS_NUEVAS','COMISIONES','VENTAS_NUEVAS'))
        self.assertFalse(incidencias(r['final'],r['r34'])[1])
        r['final']['CONCATENAR'] = ['1001','1001','1021','1021']
        casos = incidencias(r['final'],r['r34'])[1]
        self.assertEqual([c['relacionados'] for c in casos], [[1000,1001],[1002,1003]])

    def test_paso2_solo_conflicto_nuevo_inequivoco(self):
        r = resultados(('123456789001','123456789002'))
        for origen in ['COMISIONES','VENTAS_NUEVAS']:
            r['final'].loc[1,'__ORIGEN'] = origen
            for evidencia,esperado in [(['123456789003'],True), (['123456789002'],False),
                    (['1.23457E+11'],False), ([''],False),
                    (['123456789002','123456789003'],False),
                    (['123456789003','123456789004'],False)]:
                r['r34'] = pd.DataFrame({'__CONCATENAR':['1011']*len(evidencia),
                    '__SERIAL_R34':evidencia,'__ANO_R34':2026,'__MES_R34':9})
                self.assertEqual(bool(incidencias(r['final'],r['r34'])[2]), esperado and origen=='VENTAS_NUEVAS')
        r['final'].loc[1,'SERIAL'] = ''
        r['r34'] = r['r34'].iloc[:1]
        self.assertFalse(incidencias(r['final'],r['r34'])[2])

    def test_decision_serial_no_recalcula_ningun_otro_campo(self):
        r = resultados(('123456789001','123456789002'))
        d = r['final'];d['__SERIAL_COMISION'] = d['SERIAL']
        d['MONTO TX AGOSTO'] = 1500;d['MONTO TX SEPTIEMBRE'] = 0
        d['__ANO_REPORTE'] = 2026;d['__MES_REPORTE'] = 8
        d['OBSERVACION'] = 'Nota ficticia inalterable'
        r['r34'] = pd.DataFrame([{'__CONCATENAR':'1011','__SERIAL_R34':'1.23457E+11',
            'TERMINAL':'000123456789003 [001] POS','__ANO_R34':2026,'__MES_R34':9}])
        caso = incidencias(d,r['r34'])[2][0]
        self.assertEqual(caso['evidencia'],(('123456789003',(2026,9),'TERMINAL'),))
        with patch('revision_manual.recalcular_comisiones',side_effect=AssertionError('No recalcular')):
            mantener,_ = aplicar_decision(r,{},2,1001,'comisiones')
            pd.testing.assert_frame_equal(mantener['final'],d)
            nuevo,estado = aplicar_decision(r,{},2,1001,'r34',serial='123456789003')
            nuevo = preparar_descarga(nuevo,estado)
        cols = [c for c in d if c not in {'SERIAL','__SERIAL_COMISION'}]
        pd.testing.assert_frame_equal(nuevo['final'][cols],d[cols])

    def test_aplica_pago_con_periodo_y_preserva_estados_historicos(self):
        d = resultados(('A','B','C','D','E'),('COMISIONES',)*4+('VENTAS_NUEVAS',))['final']
        d['ESTATUS'] = ['PENDIENTE','PAGADO','DESINSTALADO','APLICA PAGO 02/01/2026','PENDIENTE']
        d['__ANO_REPORTE'] = 2026;d['__MES_REPORTE'] = 9
        d['MONTO TX AGOSTO'] = 1500;d['MONTO TX SEPTIEMBRE'] = 0
        d['FECHA DE PAGO'] = datetime(2026,1,2)
        d['OBSERVACION'] = 'Nota histórica ficticia'
        r34 = pd.DataFrame([{'__CONCATENAR':'1001','__MONTO_TX':0,'__ANO_R34':2026,'__MES_R34':9,'__SERIAL_R34':'A'}])
        with patch.object(p,'viernes_semana_actual',return_value=datetime(2027,2,5)):
            out = p.recalcular_comisiones(d,r34,{'100','101','102','103','104'},mes_r34=9)
        self.assertEqual(list(out['ESTATUS']),['Aplica Pago','PAGADO','DESINSTALADO','APLICA PAGO 02/01/2026','Aplica Pago'])
        self.assertEqual(out.at[0,'FECHA DE PAGO'],datetime(2027,2,5))
        pd.testing.assert_series_equal(out['OBSERVACION'],d['OBSERVACION'])
        self.assertEqual(out.at[0,'CON TX'],'CON_TX')

    def test_exportacion_respeta_esquema_y_distribuciones_historicas(self):
        d = resultados()['final'].iloc[:1].copy()
        d['ESTATUS']='PAGADO';d['BANCO']='TESORO'
        d['CANAL DE VENTA (JORNADA QUE PERTENECE)']='JORNADA DEL TESORO'
        d['VENDEDOR BANCO']='TESORO';d['VENDEDOR / FREELANCE']='Persona ficticia'
        d['VENDEDOR AGENTE AUTORIZADO']='Agente ficticio'
        d['MONTO COMISION BANCO $']=10;d['MONTO COMISION VENDEDOR/FREELANCE $']=5
        d['MONTO COMISION AGENTE AUTORIZADO $']=5;d['MONTO TOTAL A PAGAR $']=20
        d['REGISTRO DE OPERADORES ']='SI';d['OBSERVACION']='Nota histórica'
        d['MONTO TX AGOSTO']=0
        cols=[c for c in d if not c.startswith('__')]
        wb=Workbook();ws=wb.active;ws.title='VENTAS';ws.append(cols);ws.append([d.iloc[0][c] for c in cols])
        ws.auto_filter.ref=f'A1:{ws.cell(1,len(cols)).column_letter}2'
        wb.create_sheet('OTRA')['A1']='=1+1'
        source=io.BytesIO();wb.save(source)
        b=p.preparar_comisiones(source)
        r34=pd.DataFrame([{'__CONCATENAR':'1001','__MONTO_TX':2000,'__ANO_R34':2026,'__MES_R34':9,'__SERIAL_R34':'A'}])
        out=p.recalcular_comisiones(b,r34,{'100'},mes_r34=9)
        for c in cols:
            if c not in ['CON TX','MONTO TX AGOSTO']:
                self.assertEqual(out.at[0,c],b.at[0,c],c)
        raw=p.generar_excel_resultado({'final':out,'bytes_comisiones_original':source.getvalue(),'hoja_comisiones':'VENTAS','cantidad_original':1})
        result=load_workbook(io.BytesIO(raw));headers=[c.value for c in result['VENTAS'][1]]
        self.assertEqual([h for h in headers if h!='MONTO TX SEPTIEMBRE'],cols)
        self.assertEqual(headers.count('AFILIADO'),1)
        self.assertEqual(sum('OPERADORES' in str(h) for h in headers),1)
        self.assertFalse(any(str(h).startswith('__') for h in headers))
        self.assertEqual(result['OTRA']['A1'].value,'=1+1')
        self.assertIsNotNone(result['VENTAS'].auto_filter.ref)
        self.assertEqual(p.buscar_columna(out,'access'),'REGISTRO DE OPERADORES ')

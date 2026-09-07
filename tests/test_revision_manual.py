import io
import unittest
from copy import deepcopy
from datetime import date

import pandas as pd
from openpyxl import Workbook, load_workbook

from revision_manual import incidencias, aplicar_decision, preparar_descarga, normalizar_serial
from procesamiento import generar_excel_resultado
from reglas_comisiones import aplicar_motor_comisiones


def resultados(seriales=('A', 'B'), origenes=('COMISIONES', 'VENTAS_NUEVAS')):
    filas = []
    for n, (serial, origen) in enumerate(zip(seriales, origenes)):
        filas.append({'SERIAL': serial, 'CONCATENAR': str(100+n)+'1', 'AFILIADO': str(100+n),
            'TERMINAL': '1', 'FECHA': date(2026, 8, 1), 'EQUIPO': 'Castle Dynamo',
            'ESTATUS': 'Pendiente', 'ESTATUS CXC': 'COMODATO', 'CANAL': 'REGION CENTRO',
            'VENDEDOR': 'Persona ficticia', 'BANCO': 'OTRO',
            'CANAL DE VENTA (JORNADA QUE PERTENECE)': 'OFICINA', 'OBSERVACION': '', 'CON TX': 'SIN TX',
            'VENDEDOR BANCO': '', 'MONTO COMISION BANCO $': None, 'VENDEDOR / FREELANCE': '',
            'MONTO COMISION VENDEDOR/FREELANCE $': None, 'VENDEDOR AGENTE AUTORIZADO': 'REGION CENTRO',
            'MONTO COMISION AGENTE AUTORIZADO $': None, 'MONTO TOTAL A PAGAR $': None,
            '__ROW_ID': 1000+n, '__ORIGEN': origen, '__MOTIVO_REVISION': '', '__REQUIERE_REVISION': False})
    df = aplicar_motor_comisiones(pd.DataFrame(filas))
    return {'final': df, 'r34': pd.DataFrame(), 'afiliados_access': set(),
            'mes_r34': None, 'cantidad_original': origenes.count('COMISIONES')}


class RevisionTests(unittest.TestCase):
    def test_formato_conserva_prefijos_excel(self):
        from formato_revision import _serializar
        import xml.etree.ElementTree as ET
        original=b'<ns0:root xmlns:ns0="urn:root" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x14="urn:x14" mc:Ignorable="x14"/>'
        salida=_serializar(ET.fromstring(original),original)
        ET.fromstring(salida)
        self.assertIn(b'xmlns:x14="urn:x14"',salida)
    def test_cientifico_exacto_y_pinpagos_regla_compartida(self):
        from decimal import Decimal
        from procesamiento import normalizar_identificador, serial_r34_para_equipo
        for valor in ['1.23456789012E+11','1,23456789012E+11','123456789012.0',
                      Decimal('1.23456789012E+11'), 123456789012.0]:
            self.assertEqual(normalizar_identificador(valor),'123456789012')
        self.assertEqual(normalizar_identificador('1,23457E+11'),'123457000000')
        with self.assertRaises(ValueError):
            normalizar_identificador(float(2**54))
        for equipo, registro in [('Castle Dynamo', {'__SERIAL_R34':'1.23456789012E+11'}),
            ('Pinpagos', {'__SERIAL_R34':'NOCOINCIDE','TERMINAL':'000123456789012 [001] POS'}),
            ('Pinpagos', {'__SERIAL_R34':'NOCOINCIDE','__SERIAL_TERMINAL_R34':'123456789012.0'})]:
            r=resultados(('A','123456789012'));r['final'].loc[1,'EQUIPO']=equipo
            r['r34']=pd.DataFrame([dict(registro,__CONCATENAR='1011')])
            self.assertEqual(serial_r34_para_equipo(registro,equipo),'123456789012')
            self.assertFalse(incidencias(r['final'],r['r34'])[2])

    def test_defaults_no_modifican_ni_bloquean(self):
        r=resultados(('A','A'));r['final']['BANCO']='TESORO'
        r['r34']=pd.DataFrame({'__CONCATENAR':['1011'],'__SERIAL_R34':['Z']})
        self.assertTrue(all(incidencias(r['final'],r['r34']).values()))
        pd.testing.assert_frame_equal(preparar_descarga(r,{})['final'],r['final'])

    def test_elegir_historica_o_nueva_roja_exportada(self):
        from openpyxl.styles import Border, Side
        from revision_manual import tabla_comisiones
        for elegido in [1000,1001]:
            r=resultados(('A','A'));original=r['final'].copy()
            n,estado=aplicar_decision(r,{},1,1001,'desinstalado',fila_elegida=elegido)
            self.assertEqual(n['final'].loc[n['final']['__ROW_ID'].eq(elegido),'OBSERVACION'].item(),'DESINSTALADO')
            self.assertEqual(n['final'].loc[n['final']['__ROW_ID'].ne(elegido),'OBSERVACION'].item(),'')
            pd.testing.assert_series_equal(n['final']['ESTATUS'],original['ESTATUS'])
            tabla=tabla_comisiones(n['final'],[1000,1001])
            self.assertEqual(tabla.columns[tabla.columns.get_loc('ESTATUS')+1],'OBSERVACION')
            cols=[c for c in original if not c.startswith('__')]
            wb=Workbook();ws=wb.active;ws.title='VENTAS';ws.append(cols)
            ws.append([original.iloc[0][c] for c in cols])
            for c in ws[2]:
                c.border=Border(bottom=Side(style='thin'))
                c.number_format='0.00'
            b=io.BytesIO();wb.save(b)
            n.update(bytes_comisiones_original=b.getvalue(),hoja_comisiones='VENTAS')
            ws=load_workbook(io.BytesIO(generar_excel_resultado(preparar_descarga(n,estado))))['VENTAS']
            for c in ws[elegido-1000+2]:
                self.assertEqual(c.font.color.rgb,'FFFF0000')
                self.assertEqual(c.border.bottom.style,'thin')
                self.assertEqual(c.number_format,'0.00')

    def test_eliminar_historica_preserva_fila_superviviente(self):
        r=resultados(('A','B','A'),('COMISIONES','COMISIONES','VENTAS_NUEVAS'))
        original=r['final'].copy();n,estado=aplicar_decision(r,{},1,1002,'eliminar',fila_elegida=1000)
        cols=[c for c in original if not c.startswith('__')]
        wb=Workbook();ws=wb.active;ws.title='VENTAS';ws.append(cols)
        for i in range(2):ws.append([original.iloc[i][c] for c in cols])
        pos=cols.index('MONTO TOTAL A PAGAR $')+1
        ws.cell(3,pos,'=1+1')
        b=io.BytesIO();wb.save(b);n.update(bytes_comisiones_original=b.getvalue(),hoja_comisiones='VENTAS')
        ws=load_workbook(io.BytesIO(generar_excel_resultado(n)))['VENTAS']
        self.assertEqual(ws.cell(2,cols.index('SERIAL')+1).value,'B')
        self.assertEqual(ws.cell(3,cols.index('SERIAL')+1).value,'A')
        self.assertEqual(ws.cell(2,pos).value,'=1+1')
        self.assertEqual(list(n['final']['__ROW_ID']),[1001,1002])

    def test_normalizacion_y_alcance_duplicados(self):
        self.assertEqual(normalizar_serial(' a b '), 'AB')
        self.assertEqual(normalizar_serial(123.0), '123')
        for seriales, origenes, cantidad in [
            (('A','A'), ('COMISIONES','COMISIONES'), 0),
            (('A','a '), ('COMISIONES','VENTAS_NUEVAS'), 1),
            (('A','A'), ('VENTAS_NUEVAS','VENTAS_NUEVAS'), 2),
            (('A','B'), ('COMISIONES','VENTAS_NUEVAS'), 0)]:
            r = resultados(seriales, origenes)
            casos = incidencias(r['final'], r['r34'])
            self.assertEqual(len(casos[1]), cantidad)
            if cantidad:
                self.assertEqual(casos[1][0]['relacionados'], [1000, 1001])

    def test_acciones_duplicado_e_historica_protegida(self):
        for accion in ['eliminar', 'desinstalado', 'mantener']:
            r = resultados(('A', 'A')); historica = r['final'].iloc[:1].copy()
            with self.assertRaises(ValueError):
                aplicar_decision(r, {}, 1, 1000, accion)
            nuevo, estado = aplicar_decision(r, {}, 1, 1001, accion)
            pd.testing.assert_frame_equal(nuevo['final'][historica.columns].iloc[:1], historica)
            self.assertFalse(any(incidencias(nuevo['final'], nuevo['r34'], deepcopy(estado['decisiones'])).values()))
            if accion == 'eliminar':
                self.assertEqual(list(nuevo['final']['__ROW_ID']), [1000])
            elif accion == 'desinstalado':
                self.assertEqual(nuevo['final'].iloc[1]['OBSERVACION'], 'DESINSTALADO')
                self.assertEqual(nuevo['final'].iloc[1]['ESTATUS'], 'Pendiente')

    def test_r34_historico_igual_distinto_y_ambiguo(self):
        r = resultados()
        r['r34'] = pd.DataFrame({'__CONCATENAR': ['1001','1011','1011'],
                                '__SERIAL_R34': ['X','B','C']})
        casos = incidencias(r['final'], r['r34'])
        self.assertEqual(len(casos[2]), 1)
        self.assertEqual(casos[2][0]['opciones'], ('B','C'))
        with self.assertRaises(ValueError):
            aplicar_decision(r, {}, 2, 1001, 'r34')
        n, estado = aplicar_decision(r, {}, 2, 1001, 'comisiones')
        self.assertEqual(n['final'].iloc[1]['SERIAL'], 'B')
        self.assertFalse(incidencias(n['final'], n['r34'], estado['decisiones'])[2])
        n, estado = aplicar_decision(r, {}, 2, 1001, 'r34', serial='C')
        self.assertEqual(n['final'].iloc[1]['SERIAL'], 'C')
        self.assertEqual(n['final'].iloc[1]['__SERIAL_COMISION'], 'C')
        self.assertEqual(n['final'].iloc[1]['CONCATENAR'], '1011')
        self.assertFalse(incidencias(n['final'], n['r34'], estado['decisiones'])[2])
        r['r34'] = r['r34'].iloc[:2]
        self.assertFalse(incidencias(r['final'], r['r34'])[2])

    def test_serial_elegido_reabre_paso_uno(self):
        r = resultados()
        r['r34'] = pd.DataFrame({'__CONCATENAR': ['1011'], '__SERIAL_R34': ['A']})
        n, estado = aplicar_decision(r, {}, 2, 1001, 'r34', serial='A')
        self.assertEqual(len(incidencias(n['final'], n['r34'], estado['decisiones'])[1]), 1)
        preparar_descarga(n, estado)

    def test_tesoro_nuevo_mantener_o_recalcular(self):
        r = resultados()
        r['final']['BANCO'] = 'TESORO'
        casos = incidencias(r['final'], r['r34'])
        self.assertEqual([c['row_id'] for c in casos[3]], [1001])
        n, estado = aplicar_decision(r, {}, 3, 1001, 'mantener')
        pd.testing.assert_frame_equal(n['final'], r['final'])
        self.assertFalse(incidencias(n['final'], n['r34'], estado['decisiones'])[3])
        n, estado = aplicar_decision(r, {}, 3, 1001, 'jornada')
        fila = n['final'].iloc[1]
        self.assertEqual(fila['CANAL DE VENTA (JORNADA QUE PERTENECE)'], 'JORNADA BANCO DEL TESORO')
        self.assertEqual(fila['MONTO COMISION BANCO $'], 10)
        self.assertEqual(fila['MONTO COMISION AGENTE AUTORIZADO $'], 10)
        self.assertEqual(fila['MONTO TOTAL A PAGAR $'], 20)
        self.assertEqual(fila['__DIFERENCIA_CUADRE_COMISION'], 0)
        pd.testing.assert_frame_equal(n['final'].loc[:, r['final'].columns].iloc[:1], r['final'].iloc[:1])
        self.assertFalse(incidencias(n['final'], n['r34'])[3])

    def test_historicos_no_bloquean_ni_indices_definen_origen(self):
        r = resultados(('A','A'), ('COMISIONES','COMISIONES'))
        r['final']['BANCO'] = 'TESORO'
        r['final']['__REQUIERE_REVISION'] = True
        r['r34'] = pd.DataFrame({'__CONCATENAR': ['1001'], '__SERIAL_R34': ['X']})
        preparar_descarga(r, {})
        r = resultados(('A','A'))
        r['final'].index = [400, 900]
        n, estado = aplicar_decision(r, {}, 1, 1001, 'eliminar')
        self.assertEqual(list(n['final'].index), [400])

    def test_decisiones_llegan_a_xlsx_en_orden(self):
        r = resultados(('A','A','C','D'), ('COMISIONES',)+('VENTAS_NUEVAS',)*3)
        original = r['final'].copy()
        r['final'].loc[2, 'BANCO'] = 'TESORO'
        r['r34'] = pd.DataFrame({'__CONCATENAR': ['1031'], '__SERIAL_R34': ['Z']})
        pd.testing.assert_frame_equal(preparar_descarga(r, {})['final'], r['final'])
        n, estado = aplicar_decision(r, {}, 1, 1001, 'eliminar')
        n, estado = aplicar_decision(n, estado, 2, 1003, 'r34', serial='Z')
        n, estado = aplicar_decision(n, estado, 3, 1002, 'jornada')
        n = preparar_descarga(n, deepcopy(estado))
        self.assertEqual(list(n['final']['__ROW_ID']), [1000,1002,1003])
        wb = Workbook();ws=wb.active;ws.title='VENTAS'
        cols=[c for c in original if not c.startswith('__')]
        ws.append(cols);ws.append([original.iloc[0][c] for c in cols])
        wb.create_sheet('OTRA')['A1']='Conservar'
        b=io.BytesIO();wb.save(b)
        n.update(bytes_comisiones_original=b.getvalue(), hoja_comisiones='VENTAS')
        salida=load_workbook(io.BytesIO(generar_excel_resultado(n)))
        ws=salida['VENTAS']
        self.assertEqual(ws.max_row, 4)
        self.assertEqual([ws.cell(i,cols.index('SERIAL')+1).value for i in range(2,5)], ['A','C','Z'])
        self.assertEqual(ws.cell(3,cols.index('MONTO COMISION BANCO $')+1).value, 10)
        self.assertEqual(salida['OTRA']['A1'].value, 'Conservar')

    def test_desinstalado_sobrevive_revalidacion_y_exportacion(self):
        r = resultados(('A','A'))
        r['r34'] = pd.DataFrame({'__CONCATENAR': ['1011'], '__SERIAL_R34': ['Z']})
        n, estado = aplicar_decision(r, {}, 1, 1001, 'desinstalado')
        n, estado = aplicar_decision(n, estado, 2, 1001, 'r34', serial='Z')
        n = preparar_descarga(n, estado)
        self.assertEqual(n['final'].iloc[1]['OBSERVACION'], 'DESINSTALADO')
        cols=[c for c in r['final'] if not c.startswith('__')]
        wb=Workbook();ws=wb.active;ws.title='VENTAS';ws.append(cols)
        ws.append([r['final'].iloc[0][c] for c in cols])
        b=io.BytesIO();wb.save(b)
        n.update(bytes_comisiones_original=b.getvalue(),hoja_comisiones='VENTAS')
        ws=load_workbook(io.BytesIO(generar_excel_resultado(n)))['VENTAS']
        self.assertEqual(ws.cell(3,cols.index('OBSERVACION')+1).value,'DESINSTALADO')
        self.assertEqual(ws.cell(3,cols.index('SERIAL')+1).value,'Z')

    def test_decision_no_cubre_un_conflicto_diferente(self):
        r = resultados(('A','A','B'), ('COMISIONES','VENTAS_NUEVAS','VENTAS_NUEVAS'))
        r['r34'] = pd.DataFrame({'__CONCATENAR': ['1021'], '__SERIAL_R34': ['A']})
        n, estado = aplicar_decision(r, {}, 1, 1001, 'mantener')
        n, estado = aplicar_decision(n, estado, 2, 1002, 'r34', serial='A')
        self.assertEqual(len(incidencias(n['final'],n['r34'],estado['decisiones'])[1]),2)

    def test_ui_persiste_decision_entre_reruns(self):
        import sys
        import importlib.util
        from pathlib import Path
        from unittest.mock import MagicMock, patch
        fake = MagicMock()
        fake.session_state = {}
        fake.columns.return_value = [MagicMock(),MagicMock(),MagicMock()]
        fake.selectbox.side_effect = ['Marcar como DESINSTALADO',1001]
        fake.button.return_value = True
        fake.rerun.side_effect = RuntimeError('rerun')
        r=resultados(('A','A'));fake.session_state['resultados']=r
        spec=importlib.util.spec_from_file_location('ui_prueba',Path(__file__).parents[1]/'revision_ui.py')
        with patch.dict(sys.modules,{'streamlit':fake}):
            ui=importlib.util.module_from_spec(spec);spec.loader.exec_module(ui)
            with self.assertRaisesRegex(RuntimeError,'rerun'):
                ui.mostrar_revision(r,None)
            fake.button.return_value=False
            self.assertFalse(ui.mostrar_revision(fake.session_state['resultados'],None))
            self.assertIn((1,1001),fake.session_state['revision_nuevas']['decisiones'])

import unittest
from io import BytesIO
from copy import deepcopy
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from test_revision_manual import resultados
from revision_manual import incidencias, aplicar_decision, preparar_descarga
from revision_terminales import casos_terminales, preparar_cambio_terminal, aplicar_cambio_terminal
from procesamiento import generar_excel_resultado

try:
    from streamlit.testing.v1 import AppTest
except ImportError:
    AppTest = None


def ejemplo(seriales=('123456789001', '123456789001')):
    r = resultados(seriales)
    d = r['final']
    d['AFILIADO'] = '100'
    d['TERMINAL'] = ['1', '2']
    d['CONCATENAR'] = ['1001', '1002']
    d['__AFILIADO'] = '100'
    d['__TERMINAL'] = ['1', '2']
    d['__CONCATENAR'] = ['1001', '1002']
    return r


class RevisionTerminalesTests(unittest.TestCase):
    @unittest.skipIf(AppTest is None, 'Requiere Streamlit para probar la interfaz.')
    def test_interfaz_paso4_exige_preview_y_confirmacion_y_preserva_otros_pasos(self):
        app = AppTest.from_string('''
import streamlit as st
from test_revision_terminales import ejemplo
from revision_ui import mostrar_revision
if 'resultados' not in st.session_state:
    st.session_state['resultados'] = ejemplo()
pendiente = mostrar_revision(st.session_state['resultados'], {})
st.write('CAMBIO SIN APLICAR' if pendiente else 'SIN CAMBIOS SELECCIONADOS')
''', default_timeout=30).run()
        self.assertFalse(app.exception)
        self.assertTrue(any('PASO 2 DE 4' in s.value for s in app.subheader))
        self.assertTrue(any('PASO 4' in s.value for s in app.subheader))
        app.selectbox[0].select('Revisar otra terminal').run()
        app.selectbox[1].select('1').run()
        app.checkbox[0].check().run()
        self.assertEqual(len(app.session_state['resultados']['final']),2)
        app.selectbox[2].select(1000).run()
        self.assertTrue(app.button[-1].disabled)
        app.checkbox[-1].check().run()
        app.button[-1].click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.session_state['resultados']['final']),1)
        self.assertEqual(app.session_state['resultados']['final'].iloc[0]['__ORIGEN'],'COMISIONES')

    def test_mismo_serial_distinta_terminal_detectado_sin_mutar(self):
        r = ejemplo(); original = r['final'].copy(deep=True)
        casos = casos_terminales(r)
        self.assertEqual(len(casos), 1)
        self.assertTrue(casos[0]['serial_coincide'])
        self.assertEqual(casos[0]['terminales'], ['1'])
        pd.testing.assert_frame_equal(r['final'], original)

    def test_ambos_diferentes_solo_candidato_y_exige_confirmar(self):
        r = ejemplo(('123456789001', '123456789002'))
        self.assertFalse(casos_terminales(r)[0]['serial_coincide'])
        with self.assertRaisesRegex(ValueError, 'Confirme'):
            preparar_cambio_terminal(r, {}, 1001, '1')

    def test_no_mezcla_afiliados_ni_crea_caso_terminal_igual(self):
        self.assertFalse(casos_terminales(resultados()))
        r = ejemplo();r['final']['TERMINAL'] = '1'
        self.assertFalse(casos_terminales(r))

    def test_mantener_no_cambia_valores_y_oculta_caso(self):
        r=ejemplo(); vista=preparar_cambio_terminal(r,{},1001,mantener=True)
        nuevo, estado=aplicar_cambio_terminal(r,{},vista)
        pd.testing.assert_frame_equal(nuevo['final'],r['final'])
        self.assertFalse(casos_terminales(nuevo,estado))
        pd.testing.assert_frame_equal(preparar_descarga(nuevo,estado)['final'],r['final'])

    def test_colision_exige_decision_explicita_y_solo_retira_nueva(self):
        for estatus in ('PAGADO', 'DESINSTALADO', 'Pendiente'):
            r=ejemplo();r['final'].loc[0,'ESTATUS']=estatus
            original=r['final'].copy()
            with self.assertRaisesRegex(ValueError, 'ya existe'):
                preparar_cambio_terminal(r,{},1001,'1',misma_venta=True)
            vista=preparar_cambio_terminal(r,{},1001,'1',misma_venta=True,conservar_existente=1000)
            self.assertEqual(vista['eliminada'],1001)
            pd.testing.assert_frame_equal(r['final'],original)
            nuevo,estado=aplicar_cambio_terminal(r,{},vista)
            pd.testing.assert_frame_equal(nuevo['final'],original.iloc[:1])
            self.assertIn(1001,estado['eliminadas'])

    def test_no_se_puede_modificar_historica(self):
        with self.assertRaises(ValueError):
            preparar_cambio_terminal(ejemplo(),{},1000,'2',misma_venta=True)

    def test_exportacion_retiro_nueva_conserva_libro_e_historica(self):
        r=ejemplo()
        columnas=[c for c in r['final'] if not c.startswith('__')]
        w=Workbook();s=w.active;s.title='VENTAS';s.append(columnas)
        s.append([r['final'].iloc[0][c] for c in columnas])
        w.create_sheet('OTRA HOJA')['A1']='CONTENIDO FICTICIO CONSERVADO'
        b=BytesIO();w.save(b)
        r.update(bytes_comisiones_original=b.getvalue(),hoja_comisiones='VENTAS')
        vista=preparar_cambio_terminal(r,{},1001,'1',misma_venta=True,conservar_existente=1000)
        nuevo,estado=aplicar_cambio_terminal(r,{},vista)
        salida=load_workbook(BytesIO(generar_excel_resultado(preparar_descarga(nuevo,estado))))
        self.assertEqual(salida['VENTAS'].max_row,2)
        self.assertEqual(salida['VENTAS'].cell(2,columnas.index('TERMINAL')+1).value,'1')
        self.assertEqual(salida['OTRA HOJA']['A1'].value,'CONTENIDO FICTICIO CONSERVADO')

    def test_historial_tx_no_es_evidencia_de_terminal_actual(self):
        r=ejemplo()
        r['r34']=pd.DataFrame([{'__AFILIADO':'100','__TERMINAL_NUM':'3',
            '__SERIAL_R34':'123456789001','__ES_HISTORIAL_TX':True}])
        self.assertNotIn('3',casos_terminales(r)[0]['terminales'])

    def test_varias_colisiones_y_nueva_pagada_se_bloquean(self):
        r=ejemplo()
        r['final']=pd.concat([r['final'],r['final'].iloc[:1]],ignore_index=True)
        r['final'].loc[2,'__ROW_ID']=1002
        with self.assertRaises(ValueError):
            preparar_cambio_terminal(r,{},1001,'1',misma_venta=True,conservar_existente=1000)
        r=ejemplo();r['final'].loc[1,'ESTATUS']='PAGADO'
        with self.assertRaisesRegex(ValueError,'protegida'):
            preparar_cambio_terminal(r,{},1001,'1',misma_venta=True,conservar_existente=1000)

    def test_no_interpreta_texto_terminal_r34_como_numero(self):
        r=ejemplo()
        r['r34']=pd.DataFrame([{'__AFILIADO':'100','__SERIAL_R34':'123456789001',
                              'TERMINAL':'000123456789001 [003] POS'}])
        self.assertFalse(casos_terminales(r)[0]['r34'])
        r['r34']['__TERMINAL_NUM']='3'
        self.assertIn('3',casos_terminales(r)[0]['terminales'])

    def test_cambio_nuevo_con_preview_actualiza_identidad_y_tx_sin_tocar_historico(self):
        r=ejemplo();d=r['final']
        d['__ANO_REPORTE']=2026;d['__MES_REPORTE']=8
        d['__TX_VENTAS']=[{},{}]
        d['MONTO TX AGOSTO']=[555,9000]
        d.loc[1,'CON TX']='CON_TX'
        d.loc[1,'ESTATUS']='Aplica Pago'
        d['FECHA DE PAGO']=['FECHA HISTORICA','FECHA CALCULADA ANTERIOR']
        r['r34']=pd.DataFrame([{'__AFILIADO':'100','__TERMINAL_NUM':'3','__CONCATENAR':'1003',
            '__SERIAL_R34':'123456789001','__ANO_R34':2026,'__MES_R34':8,'__MONTO_TX':0}])
        original=d.copy(deep=True)
        vista=preparar_cambio_terminal(r,{},1001,'3',misma_venta=True)
        pd.testing.assert_frame_equal(r['final'],original)
        nuevo,estado=aplicar_cambio_terminal(r,{},vista)
        for c in original:
            self.assertEqual(str(nuevo['final'].loc[0,c]),str(original.loc[0,c]))
        fila=nuevo['final'].loc[1]
        self.assertEqual(fila['TERMINAL'],'3')
        self.assertEqual(fila['CONCATENAR'],'1003')
        self.assertEqual(fila['__CONCATENAR'],'1003')
        self.assertEqual(fila['MONTO TX AGOSTO'],0)
        self.assertEqual(fila['CON TX'],'SIN TX')
        self.assertEqual(fila['ESTATUS'],'Pendiente')
        self.assertTrue(pd.isna(fila['FECHA DE PAGO']))
        self.assertTrue(any(c['Campo']=='MONTO TX AGOSTO' for c in vista['cambios']))
        final=preparar_descarga(nuevo,estado)['final']
        self.assertEqual(final.loc[1,'MONTO TX AGOSTO'],0)

    def test_r34_posterior_no_traslada_monto_a_agosto(self):
        r=ejemplo();d=r['final']
        d['__ANO_REPORTE']=2026;d['__MES_REPORTE']=8
        d['__TX_VENTAS']=[{}, {(2026,8):(12,'')}]
        d['MONTO TX AGOSTO']=[33,9000]
        r['r34']=pd.DataFrame([{'__AFILIADO':'100','__TERMINAL_NUM':'3','__CONCATENAR':'1003',
            '__SERIAL_R34':'123456789001','__ANO_R34':2026,'__MES_R34':9,'__MONTO_TX':2222}])
        vista=preparar_cambio_terminal(r,{},1001,'3',misma_venta=True)
        self.assertEqual(vista['resultados']['final'].loc[1,'MONTO TX AGOSTO'],12)
        self.assertEqual(vista['resultados']['final'].loc[1,'MONTO TX SEPTIEMBRE'],2222)

    def test_vista_desactualizada_no_se_aplica(self):
        r=ejemplo();vista=preparar_cambio_terminal(r,{},1001,mantener=True)
        r['final'].loc[1,'SERIAL']='CAMBIO'
        with self.assertRaisesRegex(ValueError, 'desactualizada'):
            aplicar_cambio_terminal(r,{},vista)

    def test_paso2_actualiza_evidencia_sin_modificar_sus_reglas(self):
        r=ejemplo(('123456789001','123456789002'))
        r['r34']=pd.DataFrame([{'__CONCATENAR':'1002','__SERIAL_R34':'123456789001'}])
        antes=deepcopy(incidencias(r['final'],r['r34']))
        casos_terminales(r)
        self.assertEqual(incidencias(r['final'],r['r34']),antes)
        nuevo,estado=aplicar_decision(r,{},2,1001,'r34',serial='123456789001')
        self.assertTrue(casos_terminales(nuevo,estado)[0]['serial_coincide'])


if __name__=='__main__': unittest.main()

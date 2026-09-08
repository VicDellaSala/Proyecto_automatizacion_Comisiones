import io
import unittest
import zipfile
from unittest.mock import patch

import pandas as pd

import procesamiento as p
from openpyxl import Workbook, load_workbook


def archivo_excel(df):
    archivo = io.BytesIO()
    df.to_excel(archivo, index=False)
    archivo.seek(0)
    return archivo


class AccessTests(unittest.TestCase):
    def test_normalizacion_simetrica(self):
        for valor, esperado in [(123, "123"), (123.0, "123"), (" 123 .0 ", "123"),
                                ("00123.0", "00123"), ("00123", "00123"),
                                (pd.NA, ""), ("n / d", ""), ("AB123.0", "AB123.0")]:
            with self.subTest(valor=valor):
                self.assertEqual(p.normalizar_afiliado_access(valor), esperado)

    def test_presentes_ausentes_y_ceros_en_excel(self):
        _, afiliados = p.preparar_access(archivo_excel(pd.DataFrame({
            "AFILIADO": [101, "102.0", " 103 ", "00123", "101"],
            "CONCATENAR": ["999", "999", "999", "999", "999"],
        })))
        self.assertEqual(afiliados, {"101", "102", "103", "00123"})
        self.assertNotIn("123", afiliados)
        self.assertNotIn("999", afiliados)

    def test_marcadores_no_son_afiliados(self):
        df = pd.DataFrame({"AFILIADO": [None, pd.NA, float("nan"), "", "N/A", "N/D", "NA", "ND", "NAN", "NONE", "-", "N / D"]})
        with patch.object(p, "leer_excel_general", return_value=df):
            _, afiliados = p.preparar_access(None)
        self.assertEqual(afiliados, set())

    def test_encabezados_parciales_no_se_adivinan(self):
        for nombre in ["PRE AFILIADO", "AFILIADO ANTERIOR", "AFILIADO TERMINAL", "CODIGO", "CONCATENAR"]:
            with self.subTest(nombre=nombre):
                with self.assertRaises(ValueError):
                    p.preparar_access(archivo_excel(pd.DataFrame({nombre: [123]})))

    def test_columnas_ambiguas_y_duplicadas(self):
        for columnas in [["AFILIADO", "CODIGO_AFILIADO"], ["AFILIADO", "AFILIADO"]]:
            with self.subTest(columnas=columnas):
                with self.assertRaises(ValueError):
                    p.preparar_access(archivo_excel(pd.DataFrame([[123, 456]], columns=columnas)))

    def test_alias_exacto_existente(self):
        _, afiliados = p.preparar_access(archivo_excel(pd.DataFrame({" Código_afiliado ": [123]})))
        self.assertEqual(afiliados, {"123"})

    def test_cruce_revalidacion_y_pinpagos(self):
        base = pd.DataFrame({
            "AFILIADO": [123.0, "1234", "N/D", "00123", "999", "999"],
            "TERMINAL": [1, 1, 1, 1, 1, 2],
            "EQUIPO": ["Castle Dynamo"] * 4 + ["Pinpagos", "Zappy S1MINI2"],
            "ESTATUS": ["Pendiente"] * 6,
            "CON TX": [""] * 6,
            "SERIAL": [""] * 6,
            "ACCESS COMMERCE": ["SI"] * 6,
        })
        _, afiliados = p.preparar_access(archivo_excel(pd.DataFrame({"AFILIADO": [123, "N/D"]})))
        resultado = p.recalcular_comisiones(base, pd.DataFrame(), afiliados)
        self.assertEqual(resultado["__ACCESS_CALCULADO"].tolist(), ["SI", "NO", "NO", "NO", "NO APLICA", "NO"])
        resultado.loc[0, "AFILIADO"] = "456"
        resultado = p.recalcular_comisiones(resultado, pd.DataFrame(), afiliados)
        self.assertEqual(resultado.loc[0, "ACCESS COMMERCE"], "NO")

    def test_historicos_no_son_evidencia_del_archivo_actual(self):
        base = pd.DataFrame({
            "AFILIADO": [123, 456, 789], "TERMINAL": [1, 1, 1],
            "EQUIPO": ["Castle Dynamo", "Zappy S1MINI2", "Pinpagos"],
            "ESTATUS": ["Pagado", "Aplica Pago", "Desinstalado"],
            "ACCESS COMMERCE": ["SI", "SI", "SI"],
        })
        resultado = p.recalcular_comisiones(base, pd.DataFrame(), set())
        self.assertEqual(resultado["__ACCESS_CALCULADO"].tolist(), ["NO", "NO", "NO APLICA"])
        pd.testing.assert_frame_equal(resultado[base.columns], base)

    def test_comisiones_no_cruza_por_columna_parcial(self):
        with self.assertRaises(ValueError):
            p.recalcular_comisiones(pd.DataFrame({"PRE AFILIADO": [123]}), pd.DataFrame(), {"123"})

    def test_pipeline_y_exportacion_conservan_libro_y_filtros(self):
        libro = Workbook()
        hoja = libro.active
        hoja.title = "VENTAS"
        hoja.append(["AFILIADO", "TERMINAL", "EQUIPO", "ESTATUS", "ACCESS COMMERCE", "CON TX", "FORMULA", "MES DE CIERRE"])
        hoja.append([123, 1, "Castle Dynamo", "Pendiente", "SI", "CON_TX", "=1+1", "AGOSTO (2026)"])
        libro.create_sheet("OTRA").append(["Conservar", "=2+2"])
        maestro = io.BytesIO()
        libro.save(maestro)
        original = maestro.getvalue()
        ventas = archivo_excel(pd.DataFrame({
            "AFILIADO": [123, 456, 777, 888, 999],
            "TERMINAL": [1, 1, 0, 1, 1],
            "EQUIPO": ["Castle Dynamo", "Pinpagos", "Castle Dynamo", "SIMCARD", "Zappy S1MINI2"],
            "FECHA REPORTE": ["2026-08-01"] * 5,
            "PROPIEDAD": ["CREDICARDPOS"] * 4 + ["AGENTE AUTORIZADO"],
        }))
        # El helper escribe Sheet1; Reportes acepta COLOCACIONES.
        reporte = load_workbook(ventas)
        reporte.active.title = "COLOCACIONES"
        ventas = io.BytesIO()
        reporte.save(ventas)
        ventas.name = "ventas.xlsx"
        r34 = io.BytesIO(b"PERTENENCIA;AFIPOS;MONTO_TRANS_BS_ACUM_MES;MES_PROCESO;ANO_PROCESO\nCREDICARDPOS;1231;1;8;2026\nCREDICARDPOS;4561;1;8;2026\n")
        r34.name = "r34.csv"
        resultados = p.procesar_todo([r34], [ventas], maestro,
                                    archivo_excel(pd.DataFrame({"AFILIADO": [789]})), chunksize=1)
        self.assertEqual(resultados["columna_afiliado_access"], "AFILIADO")
        self.assertEqual(len(resultados["ventas_nuevas"]), 1)
        self.assertEqual(len(resultados["ventas_existentes"]), 1)
        self.assertEqual(len(resultados["ventas_excluidas"]), 3)
        self.assertEqual(resultados["final"]["__ACCESS_CALCULADO"].tolist(), ["NO", "NO APLICA"])
        self.assertEqual(resultados["final"]["__ESTADO_TX_CALCULADO"].tolist(), ["CON_TX", "C/P SIN TX"])
        salida = p.generar_excel_resultado(resultados)
        self.assertEqual(maestro.getvalue(), original)
        with zipfile.ZipFile(io.BytesIO(original)) as antes, zipfile.ZipFile(io.BytesIO(salida)) as despues:
            self.assertEqual(antes.namelist(), despues.namelist())
            for parte in antes.namelist():
                if parte != "xl/worksheets/sheet1.xml":
                    self.assertEqual(antes.read(parte), despues.read(parte), parte)
        exportado = load_workbook(io.BytesIO(salida))
        self.assertEqual(exportado.sheetnames, ["VENTAS", "OTRA"])
        self.assertEqual(exportado["VENTAS"]["E2"].value, "NO")
        self.assertEqual(exportado["VENTAS"]["E3"].value, "NO APLICA")
        self.assertEqual(exportado["VENTAS"]["G2"].value, "=1+1")


if __name__ == "__main__":
    unittest.main()

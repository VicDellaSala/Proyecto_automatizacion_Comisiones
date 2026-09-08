"""Inserción de columnas TX en XML sin cargar el libro completo."""
import re
import xml.etree.ElementTree as ET
from copy import deepcopy
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from formato_revision import _serializar

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def plan_columnas(final):
    anteriores = final.attrs.get('columnas_antes_tx', [])
    actuales = [c for c in final if not str(c).startswith('__')]
    return {i: actuales.index(c) + 1 for i, c in enumerate(anteriores, 1) if c in actuales}


def mover_ref(texto, mapa):
    def celda(m):
        n=column_index_from_string(m[2]);nuevo=mapa.get(n, n + max((v-k for k,v in mapa.items()), default=0))
        return m[1]+get_column_letter(nuevo)+m[3]
    texto = re.sub(r'(?<![A-Za-z0-9_])([$]?)([A-Z]{1,3})([$]?\d+)(?![A-Za-z0-9_])',celda,texto)
    def columnas(m):
        def mover(v):
            n=column_index_from_string(v.lstrip('$'));delta=max((b-a for a,b in mapa.items()),default=0)
            return ('$' if v.startswith('$') else '') + get_column_letter(mapa.get(n,n+delta))
        return mover(m[1])+':'+mover(m[2])
    return re.sub(r'(?<![A-Za-z0-9_])([$]?[A-Z]{1,3}):([$]?[A-Z]{1,3})(?![A-Za-z0-9_])', columnas, texto)


def formula(texto, mapa, local=False):
    if not texto:
        return texto
    tokens=Tokenizer('='+texto)
    for t in tokens.items:
        if t.type != 'OPERAND' or t.subtype != 'RANGE': continue
        v=t.value
        if '!' in v:
            hoja, ref=v.rsplit('!',1)
            if hoja.strip("'").upper() == 'VENTAS':
                t.value=hoja+'!'+mover_ref(ref,mapa)
        elif local:
            t.value=mover_ref(v,mapa)
    return ''.join(t.value for t in tokens.items)


def insertar_columnas_xml(xml, final):
    mapa=plan_columnas(final)
    if not any(k!=v for k,v in mapa.items()): return xml
    root=ET.fromstring(xml)
    columnas=root.find(NS+'cols')
    if columnas is not None:
        for col in list(columnas):
            columnas.remove(col)
            for n in range(int(col.get('min')), int(col.get('max'))+1):
                c=deepcopy(col); c.set('min',str(n));c.set('max',str(n));columnas.append(c)
    for el in root.iter():
        tag=el.tag.split('}')[-1]
        if tag=='c' and el.get('r'):
            el.set('r',mover_ref(el.get('r'),mapa))
        if tag=='f' and el.text:
            el.text=formula(el.text,mapa,local=True)
        if tag in {'formula','formula1','formula2'} and el.text:
            el.text=formula(el.text,mapa,local=True)
        for attr in ('ref','sqref','topLeftCell','activeCell'):
            if el.get(attr): el.set(attr,mover_ref(el.get(attr),mapa))
        if tag=='col':
            for attr in ('min','max'):
                n=int(el.get(attr));el.set(attr,str(mapa.get(n,n+max(v-k for k,v in mapa.items()))))
    return _serializar(root,xml)


def referencias_otra_parte(xml, final, local=False):
    mapa=plan_columnas(final)
    if not any(k!=v for k,v in mapa.items()): return xml
    root=ET.fromstring(xml);modificado=False
    hojas=root.find(NS+'sheets')
    local_ventas = next((str(i) for i,h in enumerate(hojas) if h.get('name','').upper()=='VENTAS'),None) if hojas is not None else None
    for el in root.iter():
        tag=el.tag.split('}')[-1]
        if tag in {'f','formula','formula1','formula2','definedName','calculatedColumnFormula','totalsRowFormula'} and el.text:
            nuevo=formula(el.text,mapa,local=local or (tag=='definedName' and local_ventas is not None and el.get('localSheetId')==local_ventas))
            if nuevo!=el.text: el.text=nuevo;modificado=True
        if local:
            for attr in ('ref','sqref'):
                if el.get(attr):
                    nuevo=mover_ref(el.get(attr),mapa)
                    if nuevo!=el.get(attr):el.set(attr,nuevo);modificado=True
    return _serializar(root,xml) if modificado else xml


def tabla_xml(xml, final):
    root=ET.fromstring(referencias_otra_parte(xml,final,local=True))
    refs=root.get('ref','').split(':')
    if len(refs)!=2: return xml
    left=column_index_from_string(re.match(r'[$]?([A-Z]+)',refs[0])[1])
    right=column_index_from_string(re.match(r'[$]?([A-Z]+)',refs[1])[1])
    cols=root.find(NS+'tableColumns')
    if cols is None: return xml
    actuales=[c for c in final if not str(c).startswith('__')]
    nuevos=final.attrs.get('columnas_tx_nuevas',{})
    for posicion,nombre in enumerate(actuales,1):
        if nombre in nuevos and left<=posicion<=right:
            col=ET.Element(NS+'tableColumn',{'id':'0','name':nombre})
            cols.insert(posicion-left,col)
    for i,col in enumerate(cols,1):col.set('id',str(i))
    cols.set('count',str(len(cols)))
    return _serializar(root,xml)


def cadena_calculo_xml(xml, final, sheet_id):
    mapa=plan_columnas(final)
    if not any(k!=v for k,v in mapa.items()):return xml
    root=ET.fromstring(xml);actual=None
    for c in root:
        actual=c.get('i',actual)
        if actual==str(sheet_id) and c.get('r'):
            c.set('r',mover_ref(c.get('r'),mapa))
    return _serializar(root,xml)

"""Separación del Excel final. No procesa ventas ni recalcula comisiones.

Las salidas contienen valores finales y estilos, no fórmulas del maestro que
puedan quedar rotas al quitar columnas/filas. El viernes solo nombra archivos.
"""
from collections import Counter
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from reglas_comisiones import normalizar_canal, normalizar_texto, requiere_access_para_venta


AGENTES = ('REGION CENTRO', 'REGION ORIENTE', 'REGION OCCIDENTE', 'CENTRO TIPO II',
           'GRANPRO', 'MULTITIENDA', 'INV TPOS', 'POSMARACAY', 'POSMGTA', 'VENEPOS', 'VIRTUALNET')
REGIONALES = ('Alexander Pacheco', 'Ezequiel Zambrano', 'Luis Garcia', 'Liseth Villanueva', 'Jorge Milanes')
BANCOS = ('TESORO', 'BANCARIBE')
GRUPOS = {'AGENTES AUTORIZADOS': AGENTES, 'REGIONALES': REGIONALES, 'BANCOS': BANCOS}
JORNADA = 'CANAL DE VENTA (JORNADA QUE PERTENECE)'
PARES = {
    'AGENTES AUTORIZADOS': ('VENDEDOR AGENTE AUTORIZADO', 'MONTO COMISION AGENTE AUTORIZADO $'),
    'REGIONALES': ('VENDEDOR / FREELANCE', 'MONTO COMISION VENDEDOR/FREELANCE $'),
    'BANCOS': ('VENDEDOR BANCO', 'MONTO COMISION BANCO $'),
}
ACCESS = ('REGISTRO DE OPERADORES', 'REGISTROS DE OPERADORES ACCESS',
          'REGISTRO DE OPERADORES ACCESS', 'REGISTRO DE OPERADORES ACCESS COMERCES',
          'REGISTRO DE OPERADORES ACCESS COMMERCE', 'ACCESS COMMERCE', 'ACCESS COMERCE')
BASE_SALIDA = ('FECHA DE PAGO', 'ESTATUS', 'OBSERVACION', 'MES DE CIERRE',
               'FECHA DE SOLICITUD RECIBIDA', 'CANAL', 'AFILIADO', 'TERMINAL',
               'FECHA', 'VENDEDOR', 'EQUIPO', 'MONTO TOTAL $', 'MONTO TOTAL BS',
               'MONTO ESTIMADO DE FACT $', 'ESQUEMA COMERCIAL', 'OPERADORA',
               'BANCO', 'RAZON SOCIAL', 'RIF', 'TLF', 'DIRECCION',
               'REPRESENTANTE LEGAL', 'CORREO', 'ESTATUS CXC', JORNADA)
COLUMNAS_REVISION = ('CANAL', 'VENDEDOR / FREELANCE', 'MONTO COMISION VENDEDOR/FREELANCE $',
                     'AFILIADO', 'TERMINAL', 'BANCO', 'EQUIPO', 'FECHA')
AVISO_REGIONAL = 'Filas con comisión regional sin destinatario inequívoco'
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


class ErrorSeparador(ValueError):
    """Errores aptos para mostrar sin datos de clientes."""


def viernes_misma_semana(hoy=None):
    """Semana lunes-domingo: sábado/domingo usan el viernes anterior."""
    hoy = hoy or datetime.now(ZoneInfo('America/Caracas')).date()
    return hoy + timedelta(days=4 - hoy.weekday())


def encabezado(valor):
    texto = normalizar_texto(valor)
    texto = re.sub(r'\.\d+$', '', texto)  # duplicados exportados por pandas
    texto = re.sub(r'\s*/\s*', '/', texto)
    texto = re.sub(r'\s*\$+', ' $', texto).strip()
    aliases = {'NUMERO DE CUENTA': 'NUM DE CUENTA', 'NRO DE CUENTA': 'NUM DE CUENTA',
               'TOTAL VENTA ESTIMADO $': 'TOTAL DE VENTAS ESTIMADO EN $',
               'TOTAL VENTAS ESTIMADO $': 'TOTAL DE VENTAS ESTIMADO EN $',
               'TOTAL DE VENTAS ESTIMADO $': 'TOTAL DE VENTAS ESTIMADO EN $'}
    return aliases.get(texto, texto)


def monto(valor):
    """Importe ya asignado; nunca infiere tarifas ni interpreta separadores ambiguos."""
    if valor is None or str(valor).strip() in ('', '-'):
        return Decimal(0)
    if isinstance(valor, bool):
        raise ErrorSeparador('Hay comisiones con importes inválidos. Revise el Excel final.')
    texto = str(valor).strip()
    if re.fullmatch(r'\d+,\d{1,2}', texto):
        texto = texto.replace(',', '.')
    try:
        n = Decimal(texto)
        if not n.is_finite() or n < 0:
            raise InvalidOperation
        return n
    except InvalidOperation:
        raise ErrorSeparador('Hay comisiones con importes inválidos. Revise el Excel final.') from None


def access_visible(fila, registros):
    if not requiere_access_para_venta(fila):
        return 'NO APLICA'
    for valor in registros:
        texto = normalizar_texto(valor)
        if texto == 'SI' or (re.fullmatch(r'\d+(?:\.0+)?', texto) and Decimal(texto) > 0):
            return 'SI'
    return 'NO'


@dataclass
class ResultadoPagos:
    contenido_zip: bytes
    nombre_zip: str
    resumen: list
    advertencias: dict
    casos_regionales_por_revisar: list = field(default_factory=list, repr=False)

    @property
    def archivos_generados(self):
        return sum(r['Registros'] > 0 for r in self.resumen)


def _columnas(celdas):
    nombres = [encabezado(c.value) for c in celdas]
    posiciones = {}
    for i, nombre in enumerate(nombres):
        posiciones.setdefault(nombre, []).append(i)
    requeridas = ['CANAL', 'EQUIPO', 'BANCO', 'VENDEDOR', JORNADA]
    requeridas += [v for par in PARES.values() for v in par]
    for nombre in requeridas:
        cantidad = len(posiciones.get(encabezado(nombre), []))
        if cantidad != 1:
            raise ErrorSeparador(f'Se necesita una columna inequívoca: {nombre}. Encontradas: {cantidad}.')
    access = [i for i, n in enumerate(nombres) if n in {encabezado(a) for a in ACCESS}]
    if not access:
        raise ErrorSeparador('Falta REGISTRO DE OPERADORES o ACCESS COMMERCE en VENTAS.')
    return nombres, posiciones, access


def _estructura_salida(posiciones, access):
    """Lista permitida: ninguna columna adicional del maestro llega al Excel."""
    seleccion = {}
    for nombre in (*BASE_SALIDA, *[n for par in PARES.values() for n in par], 'CON TX'):
        candidatas = posiciones.get(encabezado(nombre), [])
        if len(candidatas) > 1:
            if nombre in ('FECHA DE PAGO', 'OBSERVACION'):
                candidatas = candidatas[:1]  # principal inicial, no la administrativa
            elif nombre == 'AFILIADO':
                # El afiliado operativo está entre CANAL y TERMINAL, no en el
                # bloque auxiliar inicial. No se decide por el sufijo de pandas.
                terminales = posiciones.get('TERMINAL', [])
                candidatas = [i for i in candidatas if len(terminales) == 1
                              and posiciones['CANAL'][0] < i < terminales[0]]
            if len(candidatas) != 1:
                raise ErrorSeparador(f'No se puede identificar inequívocamente la columna de salida {nombre}.')
        if candidatas:
            seleccion[nombre] = candidatas[0]
    seleccion['ACCESS COMMERCE'] = access[0]
    faltantes = [n for n in (*BASE_SALIDA, 'CON TX') if n not in seleccion]
    return seleccion, faltantes


def _metadatos(datos, ruta):
    """Lee dimensiones y cachés por streaming, sin cargar las otras hojas."""
    anchos, alturas, sin_cache = [], {}, set()
    formato = None
    with ZipFile(BytesIO(datos)) as z, z.open(ruta) as xml:
        for _, e in ET.iterparse(xml, events=('end',)):
            if e.tag == NS + 'c':
                v = e.find(NS + 'v')
                if e.find(NS + 'f') is not None and (v is None or v.text is None) and e.get('t') != 'str':
                    sin_cache.add(e.get('r'))
                e.clear()
            elif e.tag == NS + 'row':
                if e.get('ht'):
                    alturas[int(e.get('r'))] = float(e.get('ht'))
                e.clear()
            elif e.tag == NS + 'col':
                anchos.append(dict(e.attrib))
                e.clear()
            elif e.tag == NS + 'sheetFormatPr':
                formato = dict(e.attrib)
    return anchos, alturas, sin_cache, formato


def _destinos(fila, avisos):
    destinos = []
    canal = normalizar_canal(fila['CANAL'])
    if canal in AGENTES:
        destinos.append(('AGENTES AUTORIZADOS', canal))
    elif canal and canal != 'CREDICARDPOS':
        avisos['Filas con CANAL sin agente reconocido'] += 1
    # Se validan los tres importes, incluso si no hay un beneficiario reconocido.
    importes = {g: monto(fila[encabezado(par[1])]) for g, par in PARES.items()}
    if importes['REGIONALES'] > 0:
        partes = {normalizar_texto(p) for p in str(fila[encabezado(PARES['REGIONALES'][0])] or '').split('/')}
        personas = [r for r in REGIONALES if normalizar_texto(r) in partes]
        if len(personas) == 1:
            destinos.append(('REGIONALES', personas[0]))
        else:
            avisos['Filas con comisión regional sin destinatario inequívoco'] += 1
    if importes['BANCOS'] > 0:
        banco = {'BANCO DEL TESORO': 'TESORO', 'TESORO': 'TESORO', 'BANCARIBE': 'BANCARIBE'}.get(
            normalizar_texto(fila['VENDEDOR BANCO']))
        if banco:
            destinos.append(('BANCOS', banco))
        else:
            avisos['Filas con comisión bancaria sin destinatario reconocido'] += 1
    return destinos


def _copiar_celda(origen, destino, estilos):
    if not hasattr(origen, '_style_id') and not hasattr(origen, 'style_id'):
        # EmptyCell: una celda ausente del XML no lleva estilo propio.
        return
    # Los índices de estilos son locales al workbook: registrar cada estilo una vez.
    clave = origen.style_id if hasattr(origen, 'style_id') else getattr(origen, '_style_id', 0)
    if clave not in estilos:
        for atributo in ('font', 'fill', 'border', 'alignment', 'protection'):
            setattr(destino, atributo, copy(getattr(origen, atributo)))
        destino.number_format = origen.number_format
        estilos[clave] = copy(destino._style)
    else:
        destino._style = estilos[clave]
    destino.value = origen.value
    # Un texto que empieza por '=' sigue siendo texto, nunca una fórmula nueva.
    if origen.data_type in ('s', 'inlineStr'):
        destino.data_type = 's'


def _excel(fuente, cabecera, filas, columnas, indice_access, indice_monto, meta):
    w = Workbook()
    w.epoch = fuente.epoch
    w.loaded_theme = fuente.loaded_theme
    s = w.active
    s.title = 'VENTAS'
    estilos = {}
    anchos, alturas, _, formato = meta
    if formato:
        for atributo in ('defaultColWidth', 'defaultRowHeight', 'baseColWidth'):
            if atributo in formato:
                setattr(s.sheet_format, atributo, float(formato[atributo]) if atributo != 'baseColWidth' else int(formato[atributo]))
    indices = [indice for _, indice in columnas]
    for nueva, (nombre, vieja) in enumerate(columnas, 1):
        _copiar_celda(cabecera[vieja], s.cell(1, nueva), estilos)
        s.cell(1, nueva, nombre)
        for ancho in anchos:
            if int(ancho['min']) <= vieja + 1 <= int(ancho['max']):
                d = s.column_dimensions[get_column_letter(nueva)]
                if 'width' in ancho:
                    d.width = float(ancho['width'])
                d.hidden = ancho.get('hidden') == '1' and vieja != indice_access
                break
    cabecera_fila = next(c.row for c in cabecera if c.value is not None)
    if cabecera_fila in alturas:
        s.row_dimensions[1].height = alturas[cabecera_fila]
    total = Decimal(0)
    for nueva_fila, (numero, celdas, access) in enumerate(filas, 2):
        for nueva_col, vieja in enumerate(indices, 1):
            destino = s.cell(nueva_fila, nueva_col)
            _copiar_celda(celdas[vieja], destino, estilos)
            if vieja == indice_access:
                destino.value = access
            elif vieja == indice_monto and celdas[vieja].value is not None and str(celdas[vieja].value).strip() not in ('', '-'):
                destino.value = monto(celdas[vieja].value)
        if numero in alturas:
            s.row_dimensions[nueva_fila].height = alturas[numero]
        total += monto(celdas[indice_monto].value)
    s.auto_filter.ref = f'A1:{get_column_letter(len(indices))}{len(filas) + 1}'
    s.freeze_panes = 'A2'
    columna_monto = indices.index(indice_monto) + 1
    columna_resumen = max(2, columna_monto)
    # La suma es un valor final: se puede leer en Excel o Streamlit sin recalcular.
    for desplazamiento, etiqueta in enumerate(('US$', 'TASA', 'BS.')):
        fila = len(filas) + 3 + desplazamiento
        rotulo = s.cell(fila, columna_resumen - 1, etiqueta)
        rotulo._style = copy(s.cell(1, columna_monto)._style)
        celda = s.cell(fila, columna_resumen, total if desplazamiento == 0 else None)
        celda._style = copy(s.cell(2, columna_monto)._style)
        celda.number_format = '#,##0.00'
    b = BytesIO()
    w.save(b)
    w.close()
    return b.getvalue()


def separar_pagos(datos, hoy=None):
    """Solo lectura; devuelve ZIP y conteos agregados, sin guardar archivos reales."""
    try:
        w = load_workbook(BytesIO(datos), read_only=True, data_only=True, keep_links=False)
    except Exception:
        raise ErrorSeparador('No se pudo leer el XLSX. Compruebe que sea un Excel válido.') from None
    try:
        if 'VENTAS' not in w.sheetnames:
            raise ErrorSeparador('El archivo debe contener la hoja VENTAS.')
        s = w['VENTAS']
        cabecera = None
        for fila in s.iter_rows(max_row=30):
            nombres = [encabezado(c.value) for c in fila]
            if 'CANAL' in nombres and 'EQUIPO' in nombres:
                ultimo = max(i for i, c in enumerate(fila) if c.value is not None)
                cabecera = fila[:ultimo + 1]
                break
        if cabecera is None:
            raise ErrorSeparador('No se encontraron encabezados CANAL y EQUIPO en VENTAS.')
        nombres, posiciones, access = _columnas(cabecera)
        seleccion, faltantes = _estructura_salida(posiciones, access)
        seriales_conservados = posiciones.get('SERIAL', [])[1:]
        columnas_grupo = {}
        for grupo in GRUPOS:
            permitidas = (*BASE_SALIDA, *PARES[grupo], 'CON TX', 'ACCESS COMMERCE')
            columnas = [(n, seleccion[n]) for n in permitidas if n in seleccion]
            # Solo se excluye la primera aparición. Una lista permite conservar
            # todos los SERIAL posteriores sin colapsar encabezados repetidos.
            for indice in seriales_conservados:
                lugar = max(((origen, i + 1) for i, (_, origen) in enumerate(columnas)
                             if origen < indice), default=(-1, 0))[1]
                columnas.insert(lugar, ('SERIAL', indice))
            columnas_grupo[grupo] = columnas
        meta = _metadatos(datos, s._worksheet_path)
        incluidos = set(seleccion.values()) | set(seriales_conservados)
        # Una fórmula sin resultado guardado no se convierte silenciosamente en un vacío.
        if any(column_index_from_string(re.match(r'[A-Z]+', c)[0]) - 1 in incluidos for c in meta[2]):
            raise ErrorSeparador('Hay fórmulas sin resultado guardado en columnas de salida. Abra y guarde el Excel con los cálculos actualizados antes de separarlo.')
        grupos = {(g, n): [] for g, lista in GRUPOS.items() for n in lista}
        avisos = Counter()
        casos_regionales = []
        if faltantes:
            avisos['Columnas solicitadas ausentes del input (no se crean): ' + ', '.join(faltantes)] = len(faltantes)
        fila_cabecera = next(c.row for c in cabecera if c.value is not None)
        for numero, celdas in enumerate(s.iter_rows(min_row=fila_cabecera + 1, max_col=len(nombres)), fila_cabecera + 1):
            fila = {n: celdas[ix[0]].value for n, ix in posiciones.items()}
            # Omite filas vacías y bloques de totales, no fechas ni estados.
            if not any(fila.get(encabezado(n)) for n in ('CANAL', *[p[0] for p in PARES.values()])):
                continue
            # La regla compartida espera los nombres canónicos con sus espacios.
            evidencia = {n: fila.get(encabezado(n)) for n in ('EQUIPO', 'BANCO', 'VENDEDOR', JORNADA)}
            visible = access_visible(evidencia, [celdas[i].value for i in access])
            pendientes_antes = avisos[AVISO_REGIONAL]
            destinos = _destinos(fila, avisos)
            if avisos[AVISO_REGIONAL] > pendientes_antes:
                casos_regionales.append({n: celdas[seleccion[n]].value for n in COLUMNAS_REVISION if n in seleccion})
            for destino in destinos:
                grupos[destino].append((numero, celdas, visible))
        fecha = viernes_misma_semana(hoy).strftime('%d-%m-%Y')
        contenido, resumen = BytesIO(), []
        with ZipFile(contenido, 'w', compression=ZIP_DEFLATED) as z:
            for (grupo, nombre), filas in grupos.items():
                resumen.append({'Grupo': grupo, 'Beneficiario': nombre, 'Registros': len(filas)})
                if filas:
                    indice_monto = posiciones[encabezado(PARES[grupo][1])][0]
                    excel = _excel(w, cabecera, filas, columnas_grupo[grupo], access[0], indice_monto, meta)
                    z.writestr(f'Pago de comisiones {nombre} {fecha}.xlsx', excel)
        return ResultadoPagos(contenido.getvalue(), f'Pagos de comisiones {fecha}.zip', resumen, dict(avisos), casos_regionales)
    finally:
        w.close()
